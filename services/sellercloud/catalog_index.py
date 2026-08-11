"""
Local catalog index built from per-brand SellerCloud search slices.

Why per-brand and not one big export: the search API caps `pageSize` at 50 and
cannot filter by MPN/ASIN, while both export endpoints REQUIRE an explicit
`ProductIds` list (there is no "export everything"). But `model.brandIds` IS a
server-side filter, and brands are small (tens to ~hundreds of products), so we
pull only the brands a vetting run touches — each slice returns full rows
(ProductID/SKU, MPN, UPC, ASIN, ShadowOf, ...). Slices are cached to disk.

Indexed by UPC / ManufacturerSKU (MPN) / ASIN / ProductID (SKU); plus brand-name->id
and per-brand / per-manufacturer prefix (mode of leading letters of existing mains,
with a confidence share because brand SKUs are often noisy/legacy).

Note: search rows carry the SKU in field `ID`; custom-export rows use `ProductID`.
`_pid()` accepts either.

CLI:  python -m services.sellercloud.catalog_index [--force] [Brand Name ...]
"""
from __future__ import annotations

import json
import os
import re
import sys
import time
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from .client import SellerCloudClient, get_sellercloud_client

SLIM_KEYS = [
    "ID", "ProductID", "UPC", "ManufacturerSKU", "ManufacturerName", "ManufacturerID",
    "BrandID", "BrandName", "ASIN", "ShadowOf", "AmazonFBASKU", "AmazonMerchantSKU",
    "CompanyID", "FulfilledBy", "QtyPerCase",
]

DATA_DIR = Path(__file__).resolve().parents[2] / "data"
SLICE_DIR = DATA_DIR / "brand_slices"


# ── helpers ──────────────────────────────────────────────────────────────
def _norm(s) -> str:
    return str(s if s is not None else "").strip()


def _key(s) -> str:
    return _norm(s).upper()


def _digits(s) -> str:
    return "".join(ch for ch in str(s or "") if ch.isdigit())


# Shadow / kit SKUs carry a channel or kit marker: '-FBA', '-FBM' (with optional
# collision digits or a company code like '-FBATRB'), or a kit tag '_QY4'.  In the
# Azure `sku_data_extended_view` many kit-shadow rows have an EMPTY ShadowOf (e.g.
# 'J&J78331_QY4-FBA'), so relying on ShadowOf alone lets a kit shadow masquerade as
# a main.  This pattern is the belt-and-suspenders check.
_SHADOW_SKU_RE = re.compile(r"_QY\d+|-FB[AM]", re.IGNORECASE)


def _is_shadow_sku(pid: str) -> bool:
    """True when a ProductID looks like a shadow/kit SKU (never a real main)."""
    return bool(_SHADOW_SKU_RE.search(pid or ""))


def _pid(r: dict) -> str:
    """SKU/ProductID from either a search row (`ID`) or an export row (`ProductID`)."""
    return _norm(r.get("ProductID") or r.get("ID"))


def _lead_prefix(pid: str) -> str:
    """Leading run of letters (and '&'): MKS16-4292->MKS, MMMT1626W->MMMT,
    ZYR206909.FDS->ZYR, P&G12345->P&G."""
    out = []
    for ch in pid:
        if ch.isalpha() or ch == "&":
            out.append(ch)
        else:
            break
    return "".join(out).upper()


def _brand_sig(brand_key: str) -> str:
    """Leading run of letters (and '&') of a brand NAME, upper-cased — the brand's
    own signature to align prefixes against.  'COLGATE'->'COLGATE',
    'PARKER LABORATORIES'->'PARKER', '3M'->'' (starts with a digit)."""
    out = []
    for ch in brand_key:
        if ch.isalpha() or ch == "&":
            out.append(ch)
        else:
            break
    return "".join(out).upper()


_ALIGN_MIN_SHARE = 0.15   # an aligned prefix must be ≥ this fraction of the mode
                          # (keeps Colgate COL/Dove DOV, excludes Cardinal's rare CAR)


def _choose_brand_prefix(brand_key: str, counter: Counter) -> tuple[str, float]:
    """Pick a brand's SKU prefix.

    Default to the outright most-used (mode) prefix. Only override it with a
    brand-letter-ALIGNED prefix when (a) the mode itself doesn't align with the
    brand name (i.e. it's likely a parent/cross-brand code like 'J&J' on Colgate,
    'ULV'=Unilever on Dove) AND (b) the aligned alternative is *competitive* with
    the mode (≥ _ALIGN_MIN_SHARE of it). This keeps Colgate→COL (COL 70 vs J&J 144)
    while NOT letting a rare aligned outlier win — e.g. Cardinal Health stays 'CAH'
    (248) instead of flipping to 'CAR' (6). Brands whose name shares no leading
    letters with any prefix ('3M'->'MMM', 'Parker'->'PRK') keep the mode.
    Returns (prefix, share-of-mains)."""
    total = sum(counter.values()) or 1
    (mode_val, mode_cnt), = counter.most_common(1)
    sig = _brand_sig(brand_key)
    mode_aligned = bool(sig) and len(mode_val) >= 2 and (
        sig.startswith(mode_val) or mode_val.startswith(sig)
    )
    if sig and not mode_aligned:
        aligned = [(p, n) for p, n in counter.items()
                   if len(p) >= 2 and (sig.startswith(p) or p.startswith(sig))]
        if aligned:
            # most-used aligned prefix; ties -> shortest (closest to the brand base)
            p, n = max(aligned, key=lambda pn: (pn[1], -len(pn[0])))
            if n >= _ALIGN_MIN_SHARE * mode_cnt:
                return p, n / total
    return mode_val, mode_cnt / total


def _slim(row: dict) -> dict:
    return {k: row.get(k) for k in SLIM_KEYS if k in row}


# ── index ────────────────────────────────────────────────────────────────
@dataclass
class CatalogIndex:
    rows: list[dict]
    by_upc: dict[str, list[dict]]
    by_mpn: dict[str, list[dict]]
    by_asin: dict[str, list[dict]]
    by_sku: dict[str, dict]
    all_skus: set[str]
    brand_to_id: dict[str, str]
    manuf_to_id: dict[str, str]
    brand_prefix: dict[str, str]
    manuf_prefix: dict[str, str]
    brand_prefix_conf: dict[str, float] = field(default_factory=dict)
    brand_to_manuf: dict[str, str] = field(default_factory=dict)
    brand_names: set = field(default_factory=set)   # all BrandName seen (upper-cased)
    built_at: float = 0.0

    def manufacturer_for_brand(self, name: str) -> str:
        return self.brand_to_manuf.get(_key(name), "")

    @staticmethod
    def _mains(rows: list[dict]) -> list[dict]:
        # A main has no ShadowOf AND no shadow/kit marker in its SKU — the second
        # check catches kit shadows ('..._QY4-FBA') that the source stores with an
        # empty ShadowOf and would otherwise be mistaken for mains.
        return [r for r in rows
                if not _norm(r.get("ShadowOf"))
                and not _is_shadow_sku(_norm(r.get("ProductID") or r.get("ID")))]

    # existence
    def main_by_upc(self, upc: str) -> list[dict]:
        return self._mains(self.by_upc.get(_digits(upc), []))

    def main_by_mpn(self, mpn: str) -> list[dict]:
        return self._mains(self.by_mpn.get(_key(mpn), []))

    def shadows_for_asin(self, asin: str) -> list[dict]:
        return self.by_asin.get(_key(asin), [])

    def sku_exists(self, sku: str) -> bool:
        return _key(sku) in self.all_skus

    def next_free_shadow(self, base_sku: str, channel: str) -> str:
        """'MMM16-4292' + 'FBA' -> first free '{base}-FBA', '-FBA2', '-FBA3', ..."""
        if not self.sku_exists(f"{base_sku}-{channel}"):
            return f"{base_sku}-{channel}"
        n = 2
        while self.sku_exists(f"{base_sku}-{channel}{n}"):
            n += 1
        return f"{base_sku}-{channel}{n}"

    # resolution
    def brand_id(self, name: str) -> str | None:
        return self.brand_to_id.get(_key(name))

    def prefix_for_brand(self, name: str) -> str | None:
        return self.brand_prefix.get(_key(name))

    def prefix_for_manufacturer(self, name: str) -> str | None:
        return self.manuf_prefix.get(_key(name))


def build_from_rows(rows: list[dict], *, built_at: float = 0.0) -> CatalogIndex:
    by_upc: dict = defaultdict(list)
    by_mpn: dict = defaultdict(list)
    by_asin: dict = defaultdict(list)
    by_sku: dict = {}
    all_skus: set = set()
    brand_to_id: dict = {}
    manuf_to_id: dict = {}
    brand_pfx: dict = defaultdict(Counter)
    manuf_pfx: dict = defaultdict(Counter)
    brand_manuf: dict = defaultdict(Counter)
    brand_names: set = set()

    for r in rows:
        pid = _pid(r)
        if not pid:
            continue
        by_sku[pid.upper()] = r
        all_skus.add(pid.upper())

        if (upc := _digits(r.get("UPC"))):
            by_upc[upc].append(r)
        if (mpn := _key(r.get("ManufacturerSKU"))):
            by_mpn[mpn].append(r)
        if (asin := _key(r.get("ASIN"))):
            by_asin[asin].append(r)

        bn, bid = _key(r.get("BrandName")), _norm(r.get("BrandID"))
        if bn:
            brand_names.add(bn)
        if bn and bid not in ("", "0"):
            brand_to_id.setdefault(bn, bid)
        mn, mid = _key(r.get("ManufacturerName")), _norm(r.get("ManufacturerID"))
        if mn and mid not in ("", "0"):
            manuf_to_id.setdefault(mn, mid)
        mn_disp = _norm(r.get("ManufacturerName"))
        if bn and mn_disp:
            brand_manuf[bn][mn_disp] += 1

        if not _norm(r.get("ShadowOf")):          # prefix learned from mains only
            if (pfx := _lead_prefix(pid)):
                if bn:
                    brand_pfx[bn][pfx] += 1
                if mn:
                    manuf_pfx[mn][pfx] += 1

    def _mode(counter: Counter) -> tuple[str, float]:
        (val, cnt), = counter.most_common(1)
        return val, cnt / sum(counter.values())

    brand_prefix, brand_conf = {}, {}
    for b, c in brand_pfx.items():
        brand_prefix[b], brand_conf[b] = _choose_brand_prefix(b, c)
    manuf_prefix = {m: _mode(c)[0] for m, c in manuf_pfx.items()}

    return CatalogIndex(
        rows=rows,
        by_upc=dict(by_upc), by_mpn=dict(by_mpn), by_asin=dict(by_asin),
        by_sku=by_sku, all_skus=all_skus,
        brand_to_id=brand_to_id, manuf_to_id=manuf_to_id,
        brand_prefix=brand_prefix, manuf_prefix=manuf_prefix,
        brand_prefix_conf=brand_conf,
        brand_to_manuf={b: c.most_common(1)[0][0] for b, c in brand_manuf.items()},
        brand_names=brand_names,
        built_at=built_at,
    )


# ── per-brand slice fetch + cache (live SellerCloud; kept for the CLI) ────────
# NOTE: SellerCloud's /api/Catalog is ~4s PER 50-item page AND serialises requests
# per account, so concurrency does not help — the live pull is inherently slow.
# The app uses the local snapshot (below) instead; this stays for the CLI / manual
# refresh only.
def load_brand_rows(
    client: SellerCloudClient,
    brand_id: int | str,
    *,
    force: bool = False,
    max_age_hours: float = 12.0,
) -> list[dict]:
    path = SLICE_DIR / f"{brand_id}.json"
    if not force and path.exists():
        try:
            blob = json.loads(path.read_text(encoding="utf-8"))
            if (time.time() - float(blob.get("built_at", 0))) / 3600.0 <= max_age_hours:
                return blob.get("rows", [])
        except Exception:
            pass
    rows = [_slim(r) for r in client.iter_catalog(page_size=50, brand_ids=[int(brand_id)])]
    SLICE_DIR.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"built_at": time.time(), "rows": rows}), encoding="utf-8")
    return rows


def build_for_brands(
    client: SellerCloudClient,
    brand_ids: list[int | str],
    *,
    force: bool = False,
    max_age_hours: float = 12.0,
) -> CatalogIndex:
    rows: list[dict] = []
    for bid in brand_ids:
        rows.extend(load_brand_rows(client, bid, force=force, max_age_hours=max_age_hours))
    return build_from_rows(rows, built_at=time.time())


# ── local snapshot (SC data.xlsx) — offline index, no SellerCloud API ────────
# The catalog lives in the user's "SC data.xlsx" (sheet 'Sku Data extended' = all
# SKUs, mostly mains; sheet 'FBA' = FBA/FBM shadow SKUs + their ASIN). We read it
# ONCE into data/sc_snapshot.json (robust to the file being locked/open in Excel
# at runtime) and build the index from that — instant, offline. Refresh by calling
# refresh_snapshot() (or `python -m services.sellercloud.catalog_index --snapshot`).
SNAPSHOT_SRC = Path(os.getenv(
    "SELLERCLOUD_SNAPSHOT_PATH",
    r"C:\Users\RaianNair\FordMed\Back Office - Documents\E-commerce Channels"
    r"\Ecom Analytics\Ryan\Offer Analytics Requests\SC data.xlsx",
))
SNAPSHOT_JSON = DATA_DIR / "sc_snapshot.json"
_LOCAL: dict = {"key": None, "index": None}


def snapshot_rows_from_xlsx(path) -> list[dict]:
    """Slim rows for build_from_rows() from the two relevant sheets."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows: list[dict] = []
    if "Sku Data extended" in wb.sheetnames:
        it = wb["Sku Data extended"].iter_rows(values_only=True)
        idx = {h: i for i, h in enumerate(next(it))}

        def g(r, n):
            i = idx.get(n)
            return r[i] if (i is not None and i < len(r)) else None

        for r in it:
            pid = _norm(g(r, "ProductID"))
            if not pid:
                continue
            rows.append({
                "ProductID": pid, "UPC": _norm(g(r, "UPC")),
                "ManufacturerSKU": _norm(g(r, "ManufacturerSKU")),
                "ManufacturerName": _norm(g(r, "Manufacturer")),
                "BrandName": _norm(g(r, "BrandName")), "ASIN": _norm(g(r, "ASIN")),
                "ShadowOf": _norm(g(r, "ShadowOf")),
                "QtyPerCase": _norm(g(r, "QtyPerCase")),
                "FulfilledBy": _norm(g(r, "FulfilledBy")),
            })
    if "FBA" in wb.sheetnames:
        it = wb["FBA"].iter_rows(values_only=True)
        idx = {h: i for i, h in enumerate(next(it))}

        def g(r, n):
            i = idx.get(n)
            return r[i] if (i is not None and i < len(r)) else None

        for r in it:
            asin = _norm(g(r, "ASIN"))
            for col, tag in (("FBA_SKU", "FBA"), ("FBM_SKU", "FBM")):
                sku = _norm(g(r, col))
                if sku:
                    rows.append({"ProductID": sku, "ASIN": asin, "ShadowOf": tag})
    return rows


def refresh_snapshot(path=None) -> int:
    """Read SC data.xlsx → write data/sc_snapshot.json. Returns the row count."""
    rows = snapshot_rows_from_xlsx(Path(path or SNAPSHOT_SRC))
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    SNAPSHOT_JSON.write_text(json.dumps({"built_at": time.time(), "rows": rows}),
                             encoding="utf-8")
    _LOCAL["key"] = None
    return len(rows)


def local_index() -> CatalogIndex:
    """CatalogIndex from the local snapshot JSON (cached in memory; rebuilt if the
    JSON changes). Builds the JSON from SNAPSHOT_SRC on first use if missing."""
    if not SNAPSHOT_JSON.exists():
        refresh_snapshot()
    key = SNAPSHOT_JSON.stat().st_mtime
    if _LOCAL["index"] is None or _LOCAL["key"] != key:
        blob = json.loads(SNAPSHOT_JSON.read_text(encoding="utf-8"))
        _LOCAL["index"] = build_from_rows(blob.get("rows", []),
                                          built_at=float(blob.get("built_at", 0.0)))
        _LOCAL["key"] = key
    return _LOCAL["index"]


# ── CLI: build + smoke-test ────────────────────────────────────────────────
def _main() -> None:
    from dotenv import load_dotenv
    load_dotenv()
    argv = sys.argv[1:]
    if "--snapshot" in argv:   # rebuild the offline snapshot from SC data.xlsx
        t0 = time.time()
        n = refresh_snapshot()
        idx = local_index()
        print(f"snapshot: {n} rows in {time.time()-t0:.1f}s -> {SNAPSHOT_JSON}")
        print(f"  brands={len(idx.brand_names)} SKUs={len(idx.all_skus)} "
              f"UPCs={len(idx.by_upc)} MPNs={len(idx.by_mpn)} ASINs={len(idx.by_asin)}")
        return
    force = "--force" in argv
    names = [a for a in argv if a != "--force"] or ["Zyrtec", "Tegaderm", "McKesson"]

    client = get_sellercloud_client()
    brands = client.get_brands()
    name_to_id = {str(b.get("Value", "")).strip().lower(): b.get("Key") for b in brands}

    ids, resolved = [], []
    for n in names:
        bid = name_to_id.get(n.strip().lower())
        if bid:
            ids.append(bid)
            resolved.append((n, bid))
    print("resolved brands:", resolved, flush=True)

    t0 = time.time()
    idx = build_for_brands(client, ids, force=force)
    print(f"\nbuilt in {time.time()-t0:.1f}s: {len(idx.rows)} rows "
          f"(mains={sum(1 for r in idx.rows if not _norm(r.get('ShadowOf')))})")
    print(f"  UPCs={len(idx.by_upc)} MPNs={len(idx.by_mpn)} ASINs={len(idx.by_asin)} SKUs={len(idx.all_skus)}")

    print("\nprefix per brand (mode, confidence):")
    for n, _ in resolved:
        pf = idx.prefix_for_brand(n)
        cf = idx.brand_prefix_conf.get(_key(n))
        print(f"   {n:14} -> {pf!r}  conf={cf:.0%}" if pf else f"   {n:14} -> (none)")

    print("\nsample existence lookups:")
    for r in idx.main_by_upc("300450206909")[:1]:
        print("   main by UPC 300450206909 ->", _pid(r))
    print("   MPN 1626W exists as main? ->", bool(idx.main_by_mpn("1626W")))
    print("   ASIN B08FMCN1GT shadows ->", [_pid(r) for r in idx.shadows_for_asin("B08FMCN1GT")])
    base = "ZYL838132"
    print(f"   next free {base}-FBA ->", idx.next_free_shadow(base, "FBA"))
    print(f"   next free {base}-FBM ->", idx.next_free_shadow(base, "FBM"))


if __name__ == "__main__":
    _main()
