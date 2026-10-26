"""
Vendor families: which vendor rows are the SAME product in another variant.
"""
from __future__ import annotations

import argparse
import json
import re
import sqlite3
import sys
from collections import Counter
from functools import lru_cache
from typing import Any

from services.analytics.matching_core.matching.attr_extract import (
    _agree,
    _tokens,
    extract_amazon,
    extract_vendor,
    size_relation,
)
from services.analytics.matching_core.matching.lexicon import colors_relation, get_lexicon, norm
from services.analytics.matching_core.matching.variation import _count_value, compare_boxes, varying_boxes

FILE_VERSION = "2026-10-22 vendor_families v3"

EXACT = "exact variant"
OTHER_VARIANT = "other variant"
NOT_STATED = "variant not stated"
UNCLEAR = "unclear"
OTHER_PRODUCT = "other product"
OTHER_BRAND = "other brand"
RELATIONS = (EXACT, OTHER_VARIANT, NOT_STATED, UNCLEAR, OTHER_PRODUCT, OTHER_BRAND)

# boxes that make a variant.  "units" = the count printed on the package (variation.compare_boxes)
VARIANT_BOXES = ("size", "color", "gender", "flavor", "scent", "shade_code", "finish", "type",
                 "volume_ml", "weight_g", "units")
# boxes for the "one side does not state it" rule (the same list as the adapter's NOT STATED rule)
UNSTATED_BOXES = ("size", "color", "gender", "flavor", "scent", "shade_code", "finish", "type", "units")
# boxes whose words are taken out of a title before the product words are read
_STRIP_BOXES = ("color", "size", "gender", "flavor", "scent", "shade_code", "tone", "finish", "type")
# boxes a family can vary by (shown in the report)
_FAMILY_BOXES = ("size", "color", "gender", "flavor", "scent", "shade_code", "tone", "finish", "type",
                 "volume_ml", "weight_g", "count")

_SIZE_WORDS = frozenset("xxs xs s m l xl xxl xxxl small medium large extra x xx xxx size sizes".split())
_UNIT_WORDS = frozenset(("oz ounce ounces fl ml l liter liters litre lb lbs pound pounds kg g gram grams mg "
                         "gal gallon count ct pack packs pk pcs pc piece pieces each ea pair pairs case box "
                         "bundle inch inches qty").split())
_LABEL = {"size": "size", "color": "color", "gender": "", "flavor": "flavor", "scent": "scent",
          "shade_code": "shade", "tone": "tone", "finish": "finish", "type": "type"}

# the main Excel export gets these columns (routers/analytics.py)
FAMILY_EXPORT_COLUMNS = ["Family", "Family name", "Vendor variant", "Listing variant", "Relation", "Relation detail"]
FAMILY_EXPORT_BLANK = [""] * len(FAMILY_EXPORT_COLUMNS)
EXPORT_MAX_CANDIDATES = 40000          # bigger runs skip the family columns (memory)


# --------------------------------------------------------------------------- #
# Product words
# --------------------------------------------------------------------------- #

@lru_cache(maxsize=4)
def _vocab(version: str) -> dict[str, frozenset[str]]:
    lex = get_lexicon()
    return {
        "colors": frozenset(w for w, (_, fam, _) in lex.colors.items() if fam != "~"),
        "genders": frozenset(lex.gender),
        "filler": lex.line_filler,
        "flavor_scent": lex.flavors | lex.scents,
        "sizes": frozenset(lex.sizes),
    }


def _v() -> dict[str, frozenset[str]]:
    return _vocab(get_lexicon().version)


def _is_variant_token(tok: str) -> bool:
    v = _v()
    return (any(c.isdigit() for c in tok) or tok in _SIZE_WORDS or tok in _UNIT_WORDS or tok in v["sizes"]
            or tok in v["colors"] or tok in v["genders"] or tok in v["flavor_scent"])


def _has_variant(text: str) -> bool:
    return any(_is_variant_token(t) for t in _tokens(text))


def strip_variant_tail(title: Any) -> str:
    """'Injinji Liner Crew (Medium, Celestine)' -> 'Injinji Liner Crew';
    'Injinji 2.0 Men's Trail Crew Toesocks, Granite, Small' -> "Injinji 2.0 Men's Trail Crew Toesocks".
    A tail is removed only when it names a size, color, gender, count ... (a known variant word)."""
    t = str(title or "").strip()
    changed = True
    while changed and t:
        changed = False
        m = re.search(r"\(([^()]*)\)\s*$", t)
        if m and _has_variant(m.group(1)):
            t, changed = t[:m.start()].rstrip(" ,-|/"), True
            continue
        parts = t.split(",")
        k = len(parts)
        while k > 1 and len(_tokens(parts[k - 1])) <= 3:
            k -= 1
        if k < len(parts) and any(_has_variant(p) for p in parts[k:]):
            t, changed = ",".join(parts[:k]).rstrip(" ,-|/"), True
    return t


def product_words(title: Any, brand: Any, fields: dict[str, dict]) -> list[str]:
    """The words that name the PRODUCT (not the variant), in title order."""
    v = _v()
    drop = set(_tokens(brand))
    for box in _STRIP_BOXES:
        entry = fields.get(box) or {}
        for value in [entry.get("value")] + [a.get("value") for a in entry.get("also") or []]:
            drop |= set(_tokens(value))
    out: list[str] = []
    for tok in _tokens(strip_variant_tail(title)):
        if len(tok) < 2 or tok in out or tok in drop or tok in v["filler"] or _is_variant_token(tok):
            continue
        out.append(tok)
    return out


def _missing_words(need: list[str], have: list[str]) -> list[str]:
    """Product words of `need` that `have` does not contain.  An abbreviation counts in both directions:
    listing 'GG' covers vendor 'Good Grips', listing 'Good Grips' covers vendor 'GG'."""
    have_set = set(have)
    covered = [w in have_set for w in need]
    for i in range(len(need)):
        for j in range(i + 2, min(i + 4, len(need)) + 1):
            if "".join(w[0] for w in need[i:j]) in have_set:
                for k in range(i, j):
                    covered[k] = True
    initials = {"".join(w[0] for w in have[i:j])
                for i in range(len(have)) for j in range(i + 2, min(i + 4, len(have)) + 1)}
    for k, w in enumerate(need):
        if not covered[k] and len(w) >= 2 and w in initials:
            covered[k] = True
    return [w for k, w in enumerate(need) if not covered[k]]


# --------------------------------------------------------------------------- #
# Brand, barcode and style-number helpers
# --------------------------------------------------------------------------- #

def _brand_key(brand: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", norm(brand))


def _same_brand(vendor_key: str, listing_key: str, listing_title: Any) -> bool:
    if not vendor_key:
        return True                                     # the vendor row names no brand: no evidence either way
    if not listing_key:                                 # the listing names no brand: look in its title
        return vendor_key in re.sub(r"[^a-z0-9]", "", norm(listing_title))
    if vendor_key == listing_key:
        return True
    short, long_ = sorted((vendor_key, listing_key), key=len)
    return len(short) >= 4 and long_.startswith(short)  # "injinji" = "injinjiinc"


def _barcode(value: Any) -> str:
    s = str(value or "").strip()
    s = re.sub(r"\.0+$", "", s)                         # 760172011852.0 (a number read from Excel)
    s = re.sub(r"\D", "", s).lstrip("0")
    return s if len(s) >= 6 else ""


def _style_matches(style: str, models: list[str]) -> bool:
    """Vendor STYLE NUMBER 052100 vs Amazon part numbers: '052100', '52100', '052100BLKS' all match.
    Leading zeros are ignored (Excel drops them); a prefix needs at least 4 characters."""
    s = (style or "").upper().lstrip("0")
    if len(s) < 3:
        return False
    for m in models:
        m2 = (m or "").upper().lstrip("0")
        if m2 == s or (len(s) >= 4 and m2.startswith(s)):
            return True
    return False


# --------------------------------------------------------------------------- #
# Variant comparison and display
# --------------------------------------------------------------------------- #

def _states(fields: dict[str, dict], box: str) -> bool:
    if box == "units":
        return _count_value(fields) is not None
    if box == "gender":                                 # "unisex" says nothing (same rule as adapter.py)
        return str((fields.get("gender") or {}).get("value") or "").lower() not in ("", "unisex")
    return bool(fields.get(box))


def _show(fields: dict[str, dict], box: str) -> str:
    if box == "units":
        n = _count_value(fields)
        return f"{n:g} count" if n else "-"
    entry = fields.get(box) or {}
    value = entry.get("value")
    if value is None:
        return "-"
    if box == "volume_ml":
        return f"{float(value):g} ml"
    if box == "weight_g":
        return f"{float(value):g} g"
    return str(value)


def variant_text(fields: dict[str, dict]) -> str:
    """'size S · color black · men'  ('?' marks a box whose own sources disagree)."""
    parts = []
    for box in ("size", "color", "gender", "flavor", "scent", "shade_code", "finish", "type",
                "volume_ml", "weight_g", "units"):
        if box == "gender":
            if not fields.get("gender"):
                continue
        elif not _states(fields, box):
            continue
        label = _LABEL.get(box, "")
        text = _show(fields, box)
        if (fields.get(box) or {}).get("conflict"):
            text += " (?)"
        parts.append(f"{label} {text}".strip())
    return " · ".join(parts)


def compare_variants(vendor: dict[str, dict], listing: dict[str, dict]) -> dict[str, list[str]]:
    """Box lists: mismatch, partial (sizes), matched, unclear, vendor_missing (only the listing states it),
    listing_missing (only the vendor row states it).  Built on variation.compare_boxes."""
    cmp = compare_boxes(vendor, listing)
    out: dict[str, list[str]] = {"mismatch": [], "partial": [], "matched": [], "unclear": [],
                                 "vendor_missing": [], "listing_missing": []}
    for box in VARIANT_BOXES:
        state = cmp.get(box)
        if state == "match":
            out["matched"].append(box)
        elif state == "mismatch":
            out["mismatch"].append(box)
        elif state == "partial":
            out["partial"].append(box)
        elif state in ("unknown", "open"):
            v_has, l_has = _states(vendor, box), _states(listing, box)
            if v_has and l_has:
                out["unclear"].append(box)
            elif l_has and box in UNSTATED_BOXES:
                out["vendor_missing"].append(box)
            elif v_has and box in UNSTATED_BOXES:
                out["listing_missing"].append(box)
    return out


def _names(boxes: list[str]) -> str:
    return ", ".join("count" if b == "units" else b.replace("_code", "") for b in boxes)


def _value(fields: dict[str, dict], box: str) -> Any:
    entry = fields.get(box) or {}
    return None if entry.get("conflict") else entry.get("value")


# --------------------------------------------------------------------------- #
# The main function (pure: no database)
# --------------------------------------------------------------------------- #

def build_families(vendor_rows: list[dict], candidates: list[dict]) -> dict[str, Any]:
    """vendor_rows: [{"row_idx", "source": {title, brand, upc, itemid, raw}, "extracted": {...AI fields}}]
    candidates:  [{"row_idx", "asin", "item": {...the Amazon item...}, ...}]
    Returns {"families": [...], "rows": {row_idx: {...}}, "listings": {asin: {...}},
             "pairs": {(row_idx, asin): {"relation", "detail", "barcode", "proof", "settle"}}}"""
    # ---- vendor rows ----
    rows: dict[Any, dict] = {}
    for vr in vendor_rows:
        src = vr.get("source") or {}
        ext = extract_vendor(src, vr.get("extracted") or {})
        fields = ext["fields"]
        brand = ext["identity"]["brand"]
        rows[vr["row_idx"]] = {
            "title": str(src.get("title") or ""),
            "brand": brand,
            "brand_key": _brand_key(brand),
            "upc": _barcode(src.get("upc")),
            "style": str((fields.get("style") or {}).get("value") or ""),
            "fields": fields,
            "words": product_words(src.get("title"), brand, fields),
            "variant": variant_text(fields),
        }

    groups: dict[tuple, list] = {}
    for idx, r in rows.items():
        key = (r["brand_key"], frozenset(r["words"])) if r["words"] else ("#row", idx)
        groups.setdefault(key, []).append(idx)
    families: list[dict] = []
    for n, (key, members) in enumerate(groups.items(), 1):
        fid = f"F{n}"
        shortest = min((rows[i]["title"] for i in members), key=len)
        brand = rows[members[0]]["brand"]
        display = strip_variant_tail(shortest) or shortest
        varying = varying_boxes([{"fields": rows[i]["fields"]} for i in members])
        families.append({
            "id": fid,
            "name": f"{brand} · {display}" if brand and norm(brand) not in norm(display) else display,
            "brand": brand,
            "brand_key": rows[members[0]]["brand_key"],
            "words": list(rows[members[0]]["words"]),
            "rows": list(members),
            "varying": {b: varying[b] for b in _FAMILY_BOXES if b in varying},
        })
        for i in members:
            rows[i]["family"] = fid
            rows[i]["family_name"] = families[-1]["name"]
    fam_by_id = {f["id"]: f for f in families}

    # rows whose vendor title names women's, per brand (sibling evidence for gender)
    women_rows: dict[str, list] = {}
    for idx, r in rows.items():
        if str((r["fields"].get("gender") or {}).get("value") or "") == "women":
            women_rows.setdefault(r["brand_key"], []).append(idx)

    # ---- listings ----
    listings: dict[str, dict] = {}
    for c in candidates:
        asin = str(c.get("asin") or "")
        if not asin or asin in listings:
            continue
        item = c.get("item") or {}
        ext = extract_amazon(item)
        fields = ext["fields"]
        brand = ext["identity"]["brand"]
        title = str(item.get("title") or "")
        model_value = str((fields.get("model") or {}).get("value") or "")
        listings[asin] = {
            "title": title,
            "brand": brand,
            "brand_key": _brand_key(brand),
            "barcodes": {b for b in (_barcode(x) for x in ext["identity"]["upcs"]) if b},
            "models": [m for m in model_value.split("/") if m],
            "fields": fields,
            "words": product_words(title, brand, fields),
            "variant": variant_text(fields),
        }

    # ---- link listing groups to vendor families ----
    lgroups: dict[tuple, list[str]] = {}
    for asin, l in listings.items():
        key = (l["brand_key"], frozenset(l["words"])) if l["words"] else ("#asin", asin)
        lgroups.setdefault(key, []).append(asin)

    def word_fit(fam: dict, words: list[str], l: dict) -> tuple[bool, str]:
        if not fam["words"] or not _same_brand(fam["brand_key"], l["brand_key"], l["title"]):
            return False, ""
        missing = _missing_words(fam["words"], words)
        if missing:
            return False, f"listing lacks the product words: {', '.join(missing)}"
        sibling_words: set[str] = set()
        for other in families:
            if other is not fam and other["brand_key"] == fam["brand_key"]:
                sibling_words |= set(other["words"])
        extra = (set(words) & sibling_words) - set(fam["words"])
        if extra:
            return False, f"listing has words of another product: {', '.join(sorted(extra))}"
        return True, ""

    def pick(hits: Counter, how: str, by_words: list[str]) -> tuple[str | None, str] | None:
        top = max(hits.values(), default=0)
        best = [fid for fid, n in hits.items() if n and n == top]
        if len(best) == 1:
            return best[0], how
        if len(best) > 1:
            both = [fid for fid in best if fid in by_words]
            return (both[0], f"{how} + words") if len(both) == 1 else (None, "ambiguous: " + ", ".join(best))
        return None

    for key, asins in lgroups.items():
        first = listings[asins[0]]
        group_codes = set().union(*(listings[a]["barcodes"] for a in asins))
        group_models = [m for a in asins for m in listings[a]["models"]]
        code_hits, style_hits = Counter(), Counter()
        for f in families:
            code_hits[f["id"]] = sum(1 for i in f["rows"] if rows[i]["upc"] and rows[i]["upc"] in group_codes)
            style_hits[f["id"]] = sum(1 for i in f["rows"]
                                      if rows[i]["style"] and _style_matches(rows[i]["style"], group_models))
        by_words = [f["id"] for f in families if word_fit(f, first["words"], first)[0]]
        chosen = pick(code_hits, "barcode", by_words) or pick(style_hits, "style number", by_words)
        if chosen is None:
            if len(by_words) == 1:
                chosen = (by_words[0], "words")
            elif len(by_words) > 1:
                most = max(len(fam_by_id[fid]["words"]) for fid in by_words)
                best = [fid for fid in by_words if len(fam_by_id[fid]["words"]) == most]
                chosen = (best[0], "words") if len(best) == 1 else (None, "ambiguous: " + ", ".join(best))
            else:
                chosen = (None, "")
        for a in asins:
            listings[a]["family"], listings[a]["link"] = chosen

    # ---- what each family's rows and listings offer (for sibling evidence) ----
    fam_listings: dict[str, list[str]] = {}
    for asin, l in listings.items():
        if l.get("family"):
            fam_listings.setdefault(l["family"], []).append(asin)
    ctx = {"rows": rows, "listings": listings, "fam_by_id": fam_by_id,
           "fam_listings": fam_listings, "women_rows": women_rows, "word_fit": word_fit}

    # ---- relation of every candidate pair ----
    pairs: dict[tuple, dict] = {}
    for c in candidates:
        idx, asin = c.get("row_idx"), str(c.get("asin") or "")
        r, l = rows.get(idx), listings.get(asin)
        if r is None or l is None or (idx, asin) in pairs:
            continue
        pairs[(idx, asin)] = _relate(idx, asin, ctx)

    return {"families": families, "rows": rows, "listings": listings, "pairs": pairs,
            "version": FILE_VERSION}


def _sibling_color(idx: Any, asin: str, ctx: dict) -> str:
    """Proof that the listing's color is another option than the row's: '' when there is none."""
    rows, listings = ctx["rows"], ctx["listings"]
    r, l = rows[idx], listings[asin]
    v, a = _value(r["fields"], "color"), _value(l["fields"], "color")
    if not v or not a:
        return ""
    fid = r.get("family")
    for other in ctx["fam_listings"].get(fid, []):
        c2 = _value(listings[other]["fields"], "color")
        if other != asin and c2 and _agree("color", v, c2):
            return f"Amazon sells {v} as its own listing ({other})"
    for i in ctx["fam_by_id"].get(fid, {}).get("rows", []):
        c2 = _value(rows[i]["fields"], "color")
        if i != idx and c2 and _agree("color", c2, a):
            return f"the vendor file has {a} as its own row ({_label(i)})"
    return ""


def _sibling_size(idx: Any, asin: str, ctx: dict) -> str:
    rows, listings = ctx["rows"], ctx["listings"]
    r, l = rows[idx], listings[asin]
    v, a = _value(r["fields"], "size"), _value(l["fields"], "size")
    if not v or not a:
        return ""
    fid = r.get("family")
    for other in ctx["fam_listings"].get(fid, []):
        s2 = _value(listings[other]["fields"], "size")
        if other != asin and s2 and size_relation(v, s2) == "same":
            return f"Amazon sells size {v} as its own listing ({other})"
    for i in ctx["fam_by_id"].get(fid, {}).get("rows", []):
        s2 = _value(rows[i]["fields"], "size")
        if i != idx and s2 and size_relation(s2, a) == "same":
            return f"the vendor file has size {a} as its own row ({_label(i)})"
    return ""


def _sibling_gender(idx: Any, asin: str, ctx: dict) -> str:
    r, l = ctx["rows"][idx], ctx["listings"][asin]
    if r["fields"].get("gender") or str(_value(l["fields"], "gender") or "") != "women":
        return ""
    # the listing's own TITLE must say it (Amazon's gender field alone is wrong too often)
    if not re.search(r"\b(women'?s?|womens|ladies)\b", l["title"], re.IGNORECASE):
        return ""
    women = [i for i in ctx["women_rows"].get(r["brand_key"], []) if i != idx]
    if not women:
        return ""
    shown = ", ".join(_label(i) for i in women[:4]) + (" ..." if len(women) > 4 else "")
    return f"the vendor file lists Women's items as their own rows ({shown}); this row is not one"


def _label(row_idx: Any) -> str:
    return f"row {row_label(row_idx)}"


def _relate(idx: Any, asin: str, ctx: dict) -> dict:
    rows, listings, fam_by_id = ctx["rows"], ctx["listings"], ctx["fam_by_id"]
    r, l = rows[idx], listings[asin]
    fam = fam_by_id.get(r.get("family"))
    out = {"relation": OTHER_PRODUCT, "detail": "", "barcode": False, "proof": None, "settle": False}
    if not _same_brand(r["brand_key"], l["brand_key"], l["title"]):
        out.update(relation=OTHER_BRAND, detail=f"listing brand '{l['brand'] or '-'}', vendor brand '{r['brand']}'")
        return out
    barcode = bool(r["upc"]) and r["upc"] in l["barcodes"]
    style_hit = bool(r["style"]) and _style_matches(r["style"], l["models"])
    out["barcode"] = barcode
    if not barcode and not style_hit and (fam is None or l.get("family") != fam["id"]):
        if l.get("family"):
            detail = f"listing belongs to family {l['family']} ({fam_by_id[l['family']]['name']})"
        elif str(l.get("link", "")).startswith("ambiguous"):
            detail = f"listing fits several families ({l['link'][11:]})"
        else:
            detail = (ctx["word_fit"](fam, l["words"], l)[1] if fam else "") or "product words differ"
        out["detail"] = detail
        return out

    cv = compare_variants(r["fields"], l["fields"])
    hard, proven, doubtful = [], [], []
    for b in cv["mismatch"]:
        if b == "color" and colors_relation(_value(r["fields"], "color"), _value(l["fields"], "color")) != "different":
            why = _sibling_color(idx, asin, ctx)
            (proven if why else doubtful).append((b, why))
        else:
            hard.append(b)
    for b in cv["partial"]:
        why = _sibling_size(idx, asin, ctx)
        (proven if why else doubtful).append((b, why))
    if "gender" in cv["vendor_missing"]:
        why = _sibling_gender(idx, asin, ctx)
        if why:
            proven.append(("gender", why))
            cv["vendor_missing"].remove("gender")

    def diff(b: str) -> str:
        return f"{_names([b])} {_show(r['fields'], b) if b != 'gender' else 'not stated'} vs {_show(l['fields'], b)}"

    if hard or proven:
        parts = [diff(b) for b in hard] + [f"{diff(b)} ({why})" for b, why in proven]
        out.update(relation=OTHER_VARIANT, proof="sibling" if proven and not hard else "variants",
                   detail="; ".join(parts) + (" (barcode matches!)" if barcode else ""),
                   settle=not barcode)
        return out
    missing_notes = []
    if cv["vendor_missing"]:
        missing_notes.append(f"vendor row does not state {_names(cv['vendor_missing'])}")
    if cv["listing_missing"]:
        missing_notes.append(f"listing does not state {_names(cv['listing_missing'])}")
    if barcode:
        out.update(relation=EXACT, proof="barcode",
                   detail="same barcode" + (f" ({'; '.join(missing_notes)})" if missing_notes else ""))
        return out
    if doubtful or cv["unclear"]:
        notes = [f"{diff(b)}: " + ("a color name the word library cannot compare, and nothing in this run "
                                   "proves they differ" if b == "color" else
                                   "a combined size, and nothing in this run proves they differ")
                 for b, _ in doubtful]
        notes += [f"{_names([b])}: one side's own sources disagree" for b in cv["unclear"]]
        out.update(relation=UNCLEAR, detail="; ".join(notes))
        return out
    if missing_notes:
        out.update(relation=NOT_STATED, detail="; ".join(missing_notes))
        return out
    if cv["matched"]:
        out.update(relation=EXACT, proof="variants", detail=f"same {_names(cv['matched'])}"
                   + (" (style number matches)" if style_hit else ""))
        return out
    varies = [b for b in (fam or {}).get("varying", {}) if b in UNSTATED_BOXES or b == "count"]
    if varies:
        out.update(relation=NOT_STATED, detail=f"the family varies by {_names(['units' if b == 'count' else b for b in varies])}; "
                                               f"neither side states it")
    else:
        out.update(relation=EXACT, detail="no variant stated on either side")
    return out


# --------------------------------------------------------------------------- #
# Database (read only)
# --------------------------------------------------------------------------- #

def row_label(row_idx: Any) -> Any:
    """The row number as the app shows it (row_idx is 0-based; the app shows row_idx + 1)."""
    try:
        return int(row_idx) + 1
    except (TypeError, ValueError):
        return row_idx


def _json(text: Any) -> dict:
    if isinstance(text, dict):
        return text
    try:
        value = json.loads(text or "{}")
    except (TypeError, ValueError):
        return {}
    return value if isinstance(value, dict) else {}


def load_run(run_id: Any, conn: sqlite3.Connection | None = None) -> tuple[list[dict], list[dict]]:
    """(vendor_rows, candidates) of one run, in the shapes build_families() takes."""
    own = conn is None
    if own:
        from services import database
        conn = database._connect()
    conn.row_factory = sqlite3.Row
    try:
        vendor_rows = [{"row_idx": r["row_idx"], "source": _json(r["data_json"]), "extracted": _json(r["extracted_json"])}
                       for r in conn.execute("SELECT row_idx, data_json, extracted_json FROM analytics_catalog_rows "
                                             "WHERE run_id = ? ORDER BY row_idx", (run_id,))]
        candidates = []
        for r in conn.execute("SELECT row_idx, asin, verdict, confidence, review_status, data_json "
                              "FROM analytics_candidates WHERE run_id = ? ORDER BY row_idx, asin", (run_id,)):
            data = _json(r["data_json"])
            candidates.append({"row_idx": r["row_idx"], "asin": r["asin"], "verdict": r["verdict"],
                               "confidence": r["confidence"], "review_status": r["review_status"],
                               "data": data, "item": data.get("amazon") or {}})
    finally:
        if own:
            conn.close()
    return vendor_rows, candidates


def export_family_columns(run_id: Any) -> dict[tuple, list]:
    """For the main Excel export: {(row_idx, asin): [family, family name, vendor variant, listing variant,
    relation, relation detail]}.  Empty (blank columns) for very big runs or when anything goes wrong:
    the export itself must never fail because of this."""
    try:
        from services import database
        conn = database._connect()
        try:
            n = conn.execute("SELECT COUNT(*) FROM analytics_candidates WHERE run_id = ?", (run_id,)).fetchone()[0]
        finally:
            conn.close()
        if n > EXPORT_MAX_CANDIDATES:
            return {}
        vendor_rows, candidates = load_run(run_id)
        result = build_families(vendor_rows, candidates)
        out: dict[tuple, list] = {}
        for key, p in result["pairs"].items():
            r = result["rows"].get(key[0]) or {}
            l = result["listings"].get(key[1]) or {}
            out[key] = [r.get("family", ""), r.get("family_name", ""), r.get("variant", ""),
                        l.get("variant", ""), p["relation"], p["detail"]]
        return out
    except Exception:  # noqa: BLE001
        return {}


# --------------------------------------------------------------------------- #
# Report: console + workbook
# --------------------------------------------------------------------------- #

_VERDICT_NAME = {"verified": "Approved", "review": "Review", "not_approved": "Not Approved"}


def pair_table(result: dict, candidates: list[dict]) -> list[dict]:
    """One line per candidate pair, ordered family -> row -> relation -> confidence."""
    order = {f["id"]: n for n, f in enumerate(result["families"])}
    lines = []
    for c in candidates:
        key = (c.get("row_idx"), str(c.get("asin") or ""))
        p = result["pairs"].get(key)
        r = result["rows"].get(key[0])
        l = result["listings"].get(key[1])
        if not p or not r or not l:
            continue
        try:
            conf = float(c.get("confidence"))
        except (TypeError, ValueError):
            conf = None
        lines.append({"Family": r["family"], "Family name": r["family_name"], "Row": row_label(key[0]),
                      "Vendor title": r["title"], "Vendor variant": r["variant"], "Source UPC": r["upc"],
                      "ASIN": key[1], "Amazon title": l["title"], "Amazon brand": l["brand"],
                      "Listing variant": l["variant"], "Listing family": l.get("family") or "",
                      "Relation": p["relation"], "Why": p["detail"],
                      "Verdict": _VERDICT_NAME.get(str(c.get("verdict")), str(c.get("verdict") or "")),
                      "Confidence": conf, "Review status": c.get("review_status") or ""})

    def row_key(value: Any) -> tuple:
        try:
            return (0, float(value), "")
        except (TypeError, ValueError):
            return (1, 0.0, str(value))
    lines.sort(key=lambda d: (order.get(d["Family"], 999), row_key(d["Row"]), RELATIONS.index(d["Relation"]),
                              -(d["Confidence"] or 0)))
    return lines


def print_report(result: dict, candidates: list[dict], out=sys.stdout) -> None:
    fams = result["families"]
    print(f"{len(result['rows'])} vendor rows -> {len(fams)} families   ({FILE_VERSION})", file=out)
    for f in fams:
        varies = "; ".join(f"{b}: {', '.join(map(str, v))}" for b, v in f["varying"].items()) or "-"
        print(f"  {f['id']:>4}  {len(f['rows']):3d} rows  {f['name']}", file=out)
        print(f"        rows {', '.join(str(row_label(i)) for i in f['rows'])}   varies by: {varies}", file=out)
        print(f"        product words: {' '.join(f['words']) or '(none)'}", file=out)
    lines = pair_table(result, candidates)
    table = Counter((d["Verdict"], d["Relation"]) for d in lines)
    verdicts = [v for v in ("Approved", "Review", "Not Approved") if any(k[0] == v for k in table)]
    verdicts += sorted({k[0] for k in table} - set(verdicts))
    print("\nRelation of each listing to its vendor row, by current verdict:", file=out)
    print(f"  {'':20s}" + "".join(f"{v:>14s}" for v in verdicts), file=out)
    for rel in RELATIONS:
        print(f"  {rel:20s}" + "".join(f"{table.get((v, rel), 0):14d}" for v in verdicts), file=out)
    settle = sum(1 for c in candidates if c.get("verdict") == "review"
                 and (result["pairs"].get((c.get("row_idx"), str(c.get("asin") or ""))) or {}).get("settle"))
    print(f"\nReview pairs the run itself proves are another variant (family_pass.py would settle them): {settle}",
          file=out)


def write_workbook(result: dict, candidates: list[dict], path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fills = {EXACT: "E2EFDA", OTHER_VARIANT: "FCE4D6", NOT_STATED: "FFF2CC", UNCLEAR: "FFF2CC",
             OTHER_PRODUCT: "EDEDED", OTHER_BRAND: "EDEDED"}
    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    wb = Workbook()
    lines = pair_table(result, candidates)

    ws = wb.active
    ws.title = "Families"
    head = ["Family", "Family name", "Rows", "Row numbers", "Varies by", "Product words"] + [r.capitalize() for r in RELATIONS]
    ws.append(head)
    counts = Counter((d["Family"], d["Relation"]) for d in lines)
    for f in result["families"]:
        ws.append([f["id"], f["name"], len(f["rows"]), ", ".join(str(row_label(i)) for i in f["rows"]),
                   "; ".join(f"{b}: {', '.join(map(str, v))}" for b, v in f["varying"].items()),
                   " ".join(f["words"])] + [counts.get((f["id"], rel), 0) for rel in RELATIONS])
    widths = [8, 45, 7, 22, 40, 35] + [14] * len(RELATIONS)

    ws2 = wb.create_sheet("Pairs")
    cols = ["Family", "Family name", "Row", "Vendor title", "Vendor variant", "Source UPC", "ASIN", "Amazon title",
            "Amazon brand", "Listing variant", "Listing family", "Relation", "Why", "Verdict", "Confidence",
            "Review status"]
    ws2.append(cols)
    for d in lines:
        ws2.append([d[c] for c in cols])
        ws2.cell(row=ws2.max_row, column=cols.index("Relation") + 1).fill = PatternFill(
            "solid", fgColor=fills.get(d["Relation"], "FFFFFF"))
    widths2 = [8, 34, 6, 34, 26, 15, 13, 60, 14, 26, 9, 18, 50, 12, 10, 14]

    for sheet, w in ((ws, widths), (ws2, widths2)):
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = bold if cell.row == 1 else font
                cell.alignment = Alignment(vertical="top", wrap_text=cell.row > 1 and cell.column in (2, 8, 13))
        for i, width in enumerate(w, 1):
            sheet.column_dimensions[get_column_letter(i)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    wb.save(path)
    return path


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Vendor families of one run (read only)")
    ap.add_argument("--run", required=True, help="run id")
    ap.add_argument("--xlsx", help="also write a review workbook to this path")
    a = ap.parse_args(argv)
    vendor_rows, candidates = load_run(a.run)
    if not vendor_rows:
        print(f"run {a.run}: no vendor rows found")
        return 1
    result = build_families(vendor_rows, candidates)
    print_report(result, candidates)
    if a.xlsx:
        print(f"\nworkbook: {write_workbook(result, candidates, a.xlsx)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())