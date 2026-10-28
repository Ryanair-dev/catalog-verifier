"""
Auto-alert: when the mailbox sync (services/vendor_offers_mailbox_sync.py)
pulls in a genuinely new offer file for one of the 5 synced vendors, check
whether that vendor is now "cheapest & unique" (the same win logic
vendor_offers_export.build_xlsx's focus_vendor column uses -- lowest cost
among every vendor carrying the UPC, ties included, sole-source counts too)
on any products, and if so build the Keepa-enriched template export scoped
to just those UPCs and email it, with the file attached, via
services/graph_mail.py send_mail() -- the exact same email mechanism the
Price Desk export's own "email it when done" option already uses (see
routers/vendor_offers.py's analytics_export completion path).

Originally built 2026-10-26 to post to Teams instead (see that date's
changelog entry in docs/features/vendor_offers.md for the Teams webhook
file-attachment investigation) -- switched to email 2026-10-27 per user
request ("revert all the changes of adding to teams. Instead send via
mail... the usual thing that worked when we used the price desk").
Recipients are fixed (see RECIPIENTS below), not configurable per-run like
the manual export's recipient field.

Scoped down with build_template_xlsx's own existing filters -- per user
request 2026-10-26, after the raw winning-UPC count turned out to be
thousands per vendor per day (effectively the vendor's whole catalog
position, not an actionable alert):
  - apply_filters=True: excludes DOG/Restricted ASINs and anything under
    the existing 50-units/month, <90% Amazon-BB-share "sellable & selling
    only" thresholds (services/vendor_offers_template_export.py).
  - only_available=True: excludes UPCs where the focus vendor itself has
    no/zero available qty right now.
  - Also only considers winning UPCs that already have an ASIN assigned
    (vo_asin_by_upc) -- a UPC with no ASIN can't be Keepa/eligibility
    enriched at all, so it would be filtered out by apply_filters anyway;
    skipping it upfront avoids wasting a Keepa batch slot on it.

Called from sync_all_from_mailbox() right before it returns -- so this runs
for both the 30-minute scheduler and the manual "sync all" endpoint, since
both go through that one function. Each vendor's work happens in its own
daemon thread (not inline) so a slow Keepa/SP-API enrichment pass never
delays the scheduler's own 30-minute cadence. The dispatch itself (this
module's own import, done lazily by the caller) also happens off the
scheduler's thread -- see vendor_offers_mailbox_sync.py's 2026-10-26
changelog entry for why: an early version of this hook, imported inline in
sync_all_from_mailbox() itself, coincided with a real live-server hang the
first time it was deployed (later root-caused to an unrelated startup
scheduler pile-up, not this module -- but the decoupling is kept as cheap
defense-in-depth regardless).

Inert (silent no-op) unless graph_mail.is_configured() (.env OUTLOOK_*) --
mirrors every other notification call site's "never break the real
feature" convention.
"""
from __future__ import annotations

import datetime as dt
import io
import logging
import threading

from openpyxl import load_workbook

from services import database
from services import graph_mail
from services import vendor_offers_export as vo_export
from services import vendor_offers_template_export as vo_tpl

log = logging.getLogger(__name__)

RECIPIENTS = ["raiann@fordmed.com", "AviS@fordmed.com", "DinaB@fordmed.com"]


def _asin_having_upcs(upcs: list[str]) -> set[str]:
    if not upcs:
        return set()
    with database._connect() as conn:
        placeholders = ",".join("?" * len(upcs))
        rows = conn.execute(
            f"SELECT DISTINCT upc FROM vo_asin_by_upc WHERE upc IN ({placeholders})", upcs
        ).fetchall()
    return {r[0] for r in rows}


def _row_count(blob: bytes) -> int:
    wb = load_workbook(io.BytesIO(blob), read_only=True)
    # The template has 4 sheets (Analytics/Order/Keepa/offer) and wb.active is
    # NOT Analytics once re-read -- counting wb.active always returned 0, so
    # the alert silently skipped every real run.
    ws = wb["Analytics"] if "Analytics" in wb.sheetnames else wb.active
    n = max(0, (ws.max_row or 1) - 1)   # minus the header row
    wb.close()
    return n


def _alert_for_vendor(vendor: str) -> None:
    try:
        winning = vo_export.winning_upcs_for_vendor(vendor)
        have_asin = _asin_having_upcs(winning)
        upcs = [u for u in winning if u in have_asin]
        if not upcs:
            log.info("[vendor_offers_auto_alert] %s: no cheapest/unique/ASIN-having UPCs, skipping", vendor)
            return
        blob = vo_tpl.build_template_xlsx(
            upcs=upcs, focus_vendor=vendor, live=True,
            apply_filters=True, only_available=True, populate_keepa=True,
        )
        count = _row_count(blob)
        if count <= 0:
            log.info("[vendor_offers_auto_alert] %s: 0 rows survived filters, skipping", vendor)
            return
        filename = f"{vendor.replace(' ', '_')}_cheapest_unique_{dt.date.today().isoformat()}.xlsx"
        subject = f"Cheapest & unique opportunity — {vendor} ({count} product{'s' if count != 1 else ''})"
        body_html = (
            f"<p><b>{vendor}</b> is now the cheapest vendor on <b>{count}</b> "
            f"product{'s' if count != 1 else ''} (sole-source or price-winning), "
            f"filtered to in-stock, non-DOG/Restricted, ≥50 units/month sellers.</p>"
            f"<p>Attached: <b>{filename}</b>.</p>"
        )
        graph_mail.send_mail(RECIPIENTS, subject, body_html, attachments=[(filename, blob)])
        log.info("[vendor_offers_auto_alert] %s: emailed opportunity alert, %d rows", vendor, count)
    except Exception:
        log.exception("[vendor_offers_auto_alert] %s: alert failed", vendor)


def process_sync_summary(summary: dict) -> None:
    """Fire an opportunity-alert thread for every vendor whose mailbox sync
    actually merged new offer rows this run (summary["results"][v]["added"] > 0
    -- the no-new-email/unchanged-since-last-sync status dicts have no
    "added" key at all, so .get(..., 0) naturally skips them)."""
    if not graph_mail.is_configured():
        return
    for vendor, result in (summary.get("results") or {}).items():
        if isinstance(result, dict) and result.get("added", 0) > 0:
            threading.Thread(
                target=_alert_for_vendor, args=(vendor,),
                name=f"vendor-offers-auto-alert-{vendor}", daemon=True,
            ).start()
