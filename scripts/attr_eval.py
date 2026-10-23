"""
Measure attribute extraction on a hand-labeled sample: rules alone vs rules + AI (the AI-expanded title).

Step 1 - make the sheet (no AI calls; uses what the runs already saved):
    uv run python scripts/attr_eval.py --make --runs 12,15,18 --n 300 --out attr_labels.xlsx
  Picks a random mix of vendor rows and Amazon listings from those runs.  For every item and every attribute it
  writes what the rules found and what rules + AI found, plus an empty TRUTH column.

Step 2 - a person fills the TRUTH columns in Excel:
    the correct value  (e.g. "black", "XS/S", "236.6 ml" or just "236.6", "3")
    "-"                 the text does not state this attribute
    =                   the RULES value is right (shortcut)
    leave it empty      skip this attribute for this item (not counted)

Step 3 - score it:
    uv run python scripts/attr_eval.py --score attr_labels.xlsx
  Per attribute: how many labeled, how often rules alone are right, how often rules + AI are right,
  how often a stated attribute is found at all (recall), and how often a found value is wrong.
"""
from __future__ import annotations

import argparse
import json
import os
import random
import re
import sqlite3
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))        # run from the repo root or scripts/

BOXES = ("color", "size", "gender", "flavor", "scent", "finish", "type", "volume_ml", "weight_g", "count")
FILE_VERSION = "2026-10-22 attr_eval v1"


def _json(text):
    try:
        value = json.loads(text or "{}")
    except (TypeError, ValueError):
        return {}
    return value if isinstance(value, dict) else {}


def _extract(kind: str, data: dict, extracted: dict | None, with_ai: bool) -> dict:
    from services.analytics.matching_core.matching import attr_extract
    os.environ["MATCH_AI_TITLE_ATTRS"] = "1" if with_ai else "0"
    ext = attr_extract.extract_vendor(data, extracted) if kind == "vendor" else attr_extract.extract_amazon(data)
    return {b: (ext["fields"].get(b) or {}).get("value") for b in BOXES}


def _text(kind: str, data: dict) -> str:
    if kind == "vendor":
        cols = [f"{k}: {v}" for k, v in (data.get("raw") or {}).items() if str(v or "").strip() and len(str(v)) <= 60]
        expanded = (data.get("ai_unabbrev") or {}).get("expanded_title")
        return (f"{data.get('title') or ''}" + (f"\n[AI-expanded: {expanded}]" if expanded else "")
                + ("\n" + "; ".join(cols[:15]) if cols else ""))
    attrs = data.get("attributes") or {}
    extra = [f"{k}: {attrs.get(k)}" for k in ("color", "size") if attrs.get(k)]
    return f"{data.get('title') or ''}" + ("\n" + "; ".join(extra) if extra else "")


def make(run_ids: list[int], n: int, out: str, seed: int = 11) -> str:
    from services import database
    conn = database._connect()
    conn.row_factory = sqlite3.Row
    items = []
    try:
        for run_id in run_ids:
            for r in conn.execute("SELECT row_idx, data_json, extracted_json FROM analytics_catalog_rows "
                                  "WHERE run_id = ?", (run_id,)).fetchall():
                items.append(("vendor", run_id, f"row {int(r['row_idx']) + 1}", _json(r["data_json"]),
                              _json(r["extracted_json"])))
            seen = set()
            for r in conn.execute("SELECT asin, json_extract(data_json, '$.amazon') AS amazon FROM analytics_candidates "
                                  "WHERE run_id = ?", (run_id,)).fetchall():
                if r["asin"] not in seen:
                    seen.add(r["asin"])
                    items.append(("listing", run_id, r["asin"], _json(r["amazon"]), None))
    finally:
        conn.close()
    rng = random.Random(seed)
    vendor = [i for i in items if i[0] == "vendor"]
    listing = [i for i in items if i[0] == "listing"]
    half = n // 2
    chosen = rng.sample(vendor, min(half, len(vendor))) + rng.sample(listing, min(n - half, len(listing)))

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Labels"
    head = ["item", "run", "kind", "text"]
    for b in BOXES:
        head += [f"{b} rules", f"{b} rules+AI", f"{b} TRUTH"]
    ws.append(head)
    for kind, run_id, label, data, extracted in chosen:
        rules = _extract(kind, data, extracted, with_ai=False)
        with_ai = _extract(kind, data, extracted, with_ai=True)
        row = [label, run_id, kind, _text(kind, data)]
        for b in BOXES:
            row += [rules[b] if rules[b] is not None else "", with_ai[b] if with_ai[b] is not None else "", ""]
        ws.append(row)
    truth_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == 4)
            if ws.cell(row=1, column=cell.column).value.endswith("TRUTH"):
                cell.fill = truth_fill
    ws.column_dimensions["D"].width = 70
    ws.freeze_panes = "E2"
    help_ws = wb.create_sheet("How to fill")
    for line in ("Fill the yellow TRUTH columns:", "  the correct value (black, XS/S, 236.6, 3 ...)",
                 "  -   the text does not state this attribute", "  =   the RULES value is right",
                 "  empty: skip (not counted)", "", f"made by {FILE_VERSION}"):
        help_ws.append([line])
    wb.save(out)
    return out


def _same(box: str, a, b) -> bool:
    a, b = ("" if a is None else str(a)).strip().lower(), ("" if b is None else str(b)).strip().lower()
    if a in ("", "-") or b in ("", "-"):
        return a in ("", "-") and b in ("", "-")
    if box in ("volume_ml", "weight_g", "count"):
        try:
            x, y = float(re.sub(r"[^\d.]", "", a)), float(re.sub(r"[^\d.]", "", b))
            return abs(x - y) <= 0.02 * max(abs(x), abs(y), 1e-9)
        except ValueError:
            return False
    norm = lambda v: "/".join(sorted(p.strip() for p in v.replace(" ", "").split("/") if p.strip()))
    return norm(a) == norm(b) or a.replace("grey", "gray") == b.replace("grey", "gray")


def score(path: str) -> int:
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True)["Labels"]
    rows = list(ws.iter_rows(values_only=True))
    head = list(rows[0])
    stats = {b: {"n": 0, "rules_ok": 0, "ai_ok": 0, "stated": 0, "rules_found": 0, "ai_found": 0,
                 "rules_wrong": 0, "ai_wrong": 0} for b in BOXES}
    for row in rows[1:]:
        for b in BOXES:
            truth = row[head.index(f"{b} TRUTH")]
            if truth is None or str(truth).strip() == "":
                continue
            rules, ai = row[head.index(f"{b} rules")], row[head.index(f"{b} rules+AI")]
            truth = rules if str(truth).strip() == "=" else truth
            s = stats[b]
            s["n"] += 1
            s["rules_ok"] += _same(b, rules, truth)
            s["ai_ok"] += _same(b, ai, truth)
            stated = str(truth).strip() not in ("", "-")
            if stated:
                s["stated"] += 1
                s["rules_found"] += bool(str(rules or "").strip())
                s["ai_found"] += bool(str(ai or "").strip())
            s["rules_wrong"] += bool(str(rules or "").strip()) and not _same(b, rules, truth)
            s["ai_wrong"] += bool(str(ai or "").strip()) and not _same(b, ai, truth)
    pct = lambda a, b: f"{100 * a / b:5.1f}%" if b else "    -"
    print(f"{'attribute':10s} {'labeled':>7s} | {'right: rules':>12s} {'rules+AI':>9s} | "
          f"{'found when stated: rules':>24s} {'rules+AI':>9s} | {'wrong values: rules':>19s} {'rules+AI':>9s}")
    for b, s in stats.items():
        if not s["n"]:
            continue
        print(f"{b:10s} {s['n']:>7d} | {pct(s['rules_ok'], s['n']):>12s} {pct(s['ai_ok'], s['n']):>9s} | "
              f"{pct(s['rules_found'], s['stated']):>24s} {pct(s['ai_found'], s['stated']):>9s} | "
              f"{s['rules_wrong']:>19d} {s['ai_wrong']:>9d}")
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Attribute extraction eval: rules vs rules + AI")
    ap.add_argument("--make", action="store_true")
    ap.add_argument("--runs", default="", help="comma-separated run ids (with --make)")
    ap.add_argument("--n", type=int, default=300)
    ap.add_argument("--out", default="attr_labels.xlsx")
    ap.add_argument("--score", default="", help="a filled labels workbook")
    a = ap.parse_args(argv)
    if a.make:
        runs = [int(x) for x in a.runs.split(",") if x.strip()]
        if not runs:
            ap.error("--make needs --runs")
        print("written:", make(runs, a.n, a.out))
    if a.score:
        return score(a.score)
    return 0


if __name__ == "__main__":
    sys.exit(main())