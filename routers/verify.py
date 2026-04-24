"""
Verification engine endpoints.

Flow
----
1. Client uploads the catalog Excel and Amazon Excel (+ source type +
   an ``ai_mode`` flag).
2. Server parses both, persists Amazon rows to ``keepa_imports`` or
   ``amazon_imports`` so they can be reused across sessions, joins on ASIN,
   runs the confidence engine (using the thresholds stored in ``settings``),
   flags duplicate Item IDs, consults the blacklist, and caches extracted
   attributes for each (UPC, ASIN) pair.
3. Response carries per-row scoring plus a ``review_status`` field the
   frontend mutates during override.
"""
from __future__ import annotations

import io
import json as _json
from typing import Any, Optional

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel

from services import database
from services.confidence import score_row
from services.extractor import ai_extract, rule_extract

router = APIRouter()


# --------------------------------------------------------------------------- #
# Column maps (unchanged)
# --------------------------------------------------------------------------- #

CATALOG_COLUMNS = ["UPC/EAN", "Item ID", "Vendor Title", "Brand", "ASIN"]


def _workbook_to_rows(data: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        return []
    result: list[dict] = []
    for row in rows_iter:
        if row is None or all(v in (None, "") for v in row):
            continue
        record = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        result.append(record)
    return result


def _index_amazon(rows: list[dict], source: str) -> dict[str, dict]:
    key = "asin" if source == "amazon" else "ASIN"
    out: dict[str, dict] = {}
    for row in rows:
        asin = row.get(key)
        if asin:
            out[str(asin).strip().upper()] = row
    return out


def _find_duplicates(rows: list[dict], column: str) -> set[Any]:
    seen: dict[Any, int] = {}
    for row in rows:
        value = row.get(column)
        if value in (None, ""):
            continue
        key = str(value).strip()
        seen[key] = seen.get(key, 0) + 1
    return {k for k, v in seen.items() if v > 1}


def _failed_signals(signals: dict) -> list[str]:
    labels = {
        "upc": "UPC", "item_id": "Item ID", "brand": "Brand",
        "title": "Title", "pack": "Pack",
    }
    return [labels[k] for k, v in signals.items()
            if isinstance(v, dict) and not v.get("matched")]


# --------------------------------------------------------------------------- #
# Main verification endpoint
# --------------------------------------------------------------------------- #

@router.post("/verify")
async def verify(
    catalog_file: UploadFile = File(...),
    amazon_file: UploadFile = File(...),
    amazon_source: str = Form(...),
    ai_mode: str = Form("false"),
    abbreviations: Optional[str] = Form(None),
) -> dict:
    """Run the full verification pipeline on two uploaded Excel files."""
    if amazon_source not in ("keepa", "amazon"):
        raise HTTPException(status_code=400, detail="amazon_source must be 'keepa' or 'amazon'")

    use_ai = str(ai_mode).lower() in ("1", "true", "yes", "on")

    # Prefer the categorised library from SQLite; the form-posted list is a
    # fallback for older clients.
    abbr_list = database.flat_library()
    if not abbr_list and abbreviations:
        try:
            abbr_list = _json.loads(abbreviations) or []
        except _json.JSONDecodeError:
            abbr_list = []

    thresholds = database.get_thresholds()

    catalog_bytes = await catalog_file.read()
    amazon_bytes = await amazon_file.read()

    try:
        catalog_rows = _workbook_to_rows(catalog_bytes)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Catalog parse failed: {exc}")
    try:
        amazon_rows = _workbook_to_rows(amazon_bytes)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Amazon parse failed: {exc}")

    # Persist Amazon rows for future sessions.
    target_table = "keepa_imports" if amazon_source == "keepa" else "amazon_imports"
    try:
        database.save_imports(target_table, amazon_rows)
    except Exception:  # noqa: BLE001 — persistence must never fail the request
        pass

    amazon_index = _index_amazon(amazon_rows, amazon_source)
    dupes = _find_duplicates(catalog_rows, "Item ID")

    results = []
    for row in catalog_rows:
        asin = str(row.get("ASIN") or "").strip().upper()
        if not asin:
            continue
        upc  = str(row.get("UPC/EAN") or "").strip()
        amz = amazon_index.get(asin)

        score = score_row(row, amz, abbr_list, thresholds=thresholds)

        # Cache normalised attributes so the Pair Manager / UI can replay them.
        try:
            attrs = rule_extract(row.get("Vendor Title") or "", abbr_list)
            if use_ai:
                ai_result = ai_extract(
                    row.get("Vendor Title") or "",
                    extra_context=_safe_amz_context(amz),
                )
                if isinstance(ai_result, dict) and "error" not in ai_result:
                    attrs.update({k: v for k, v in ai_result.items() if v})
            if upc and asin:
                database.cache_attributes(
                    upc, asin, attrs,
                    source="ai+rule" if use_ai else "rule",
                )
        except Exception:  # noqa: BLE001
            pass

        results.append({
            "UPC/EAN": row.get("UPC/EAN"),
            "Item ID": row.get("Item ID"),
            "Vendor Title": row.get("Vendor Title"),
            "Brand": row.get("Brand"),
            "ASIN": row.get("ASIN"),
            "confidence": score["confidence"],
            "verdict": score["verdict"],
            "original_verdict": score["verdict"],
            "review_status": "",   # Reviewed / Manually Approved / Manually Rejected
            "signals": score["signals"],
            "amz_pack": score["amz_pack"],
            "barcode_db": None,
            "duplicate": str(row.get("Item ID") or "").strip() in dupes,
            "blacklisted": database.is_blacklisted(upc, asin),
            "notes": score["notes"],
        })

    return {
        "source": amazon_source,
        "catalog_count": len(catalog_rows),
        "amazon_count": len(amazon_rows),
        "duplicate_item_ids": sorted(dupes),
        "thresholds": thresholds,
        "ai_mode": use_ai,
        "results": results,
    }


def _safe_amz_context(amz: dict | None) -> str:
    if not amz:
        return ""
    parts = []
    for k in ("Title", "item_name", "Brand", "brand", "Size", "size"):
        v = amz.get(k)
        if v:
            parts.append(f"{k}: {v}")
    return " | ".join(parts)[:800]


# --------------------------------------------------------------------------- #
# AI re-score (single row) + attribute cache endpoints
# --------------------------------------------------------------------------- #

class AIScoreRequest(BaseModel):
    title: str
    context: str = ""


@router.post("/verify/ai")
async def verify_ai(req: AIScoreRequest) -> dict:
    extracted = ai_extract(req.title, req.context)
    # If the AI returned unabbreviated terms we don't already know about,
    # auto-add them to "Product Attributes". Silent best-effort.
    try:
        if isinstance(extracted, dict):
            for key in ("form", "product_type"):
                value = extracted.get(key)
                if isinstance(value, str) and value.strip() \
                        and not database.has_library_entry(value):
                    database.add_library_entry("Product Attributes", value, value)
    except Exception:  # noqa: BLE001
        pass
    return {"extracted": extracted}


class CacheKey(BaseModel):
    upc: str
    asin: str


@router.get("/attributes/{upc}/{asin}")
async def get_attributes(upc: str, asin: str) -> dict:
    cached = database.get_cached_attributes(upc, asin)
    return {"cached": cached}


@router.post("/attributes/clear")
async def clear_attributes(body: CacheKey) -> dict:
    """Clear the cached attributes for a row. The frontend immediately follows
    up with a barcode lookup request to repopulate the cache."""
    deleted = database.clear_cached_attributes(body.upc, body.asin)
    return {"cleared": deleted > 0}


# --------------------------------------------------------------------------- #
# Blacklist maintenance from the verify flow
# --------------------------------------------------------------------------- #

class BlacklistRequest(BaseModel):
    upc: str
    asin: str
    confidence: float | None = None
    failed_signals: list[str] = []


@router.post("/blacklist")
async def blacklist_pair(body: BlacklistRequest) -> dict:
    database.add_to_blacklist(
        body.upc, body.asin, body.confidence, body.failed_signals,
    )
    return {"ok": True}


@router.post("/verified")
async def mark_verified(body: dict) -> dict:
    upc = body.get("upc") or ""
    asin = body.get("asin") or ""
    review_status = body.get("review_status") or ""
    database.save_verified(upc, asin, body, review_status=review_status)
    return {"ok": True}
