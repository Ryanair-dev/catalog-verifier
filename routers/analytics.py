"""
Analytics tab (ROI & Cost) — FastAPI router.

Endpoints
---------
POST /api/analytics/preview  — upload file, return first N raw rows so the
                              wizard can show the header-row picker.
POST /api/analytics/runs     — create a run from the wizard payload + file,
                              kick off the 3-tier SP-API search + vetting
                              in a background thread.
GET  /api/analytics/runs     — list runs (newest first) for the history
                              table in the Analytics tab.
GET  /api/analytics/runs/{id}
                             — single run detail (progress + counts + the
                              first 200 candidates so the UI can render
                              a result grid without a second round-trip).
GET  /api/analytics/status   — shallow SP-API credentials probe.

The heavy lifting lives in `services.analytics.runner` — this router is
just the HTTP glue.
"""
from __future__ import annotations

import io
import json
import sqlite3
from typing import Any

from fastapi import APIRouter, Body, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.cell import WriteOnlyCell

from services import database
from services.file_parser import parse_raw_rows as _parse_raw_rows_shared, get_sheet_names as _get_sheet_names
from services.safety import clamp_int, read_upload_limited, safe_spreadsheet_row
from services.analytics import parse_source_rows, start_analytics_run
from services.analytics.runner import (
    request_pause, request_stop, resume_run, start_rescore,
    _tier1_upc, _tier2_itemid, _tier3_title,
    normalize_amazon_item,
    AUTO_APPROVE, REVIEW_FLOOR, MIN_CONFIDENCE,
    _is_media_format,
)
from services.analytics.matcher import calculate_confidence
from services.analytics.parser import SourceRow as _SourceRow
from services.analytics.product_categorizer import categorize, category_distance, UNKNOWN, MEDICAL
from services.analytics.ai_check import start_ai_check, stop_ai_check, is_running as ai_check_running, estimate_cost as ai_check_cost
from services.analytics.eligibility_check import (
    start_eligibility_check, stop_eligibility_check,
    is_running as elig_check_running, approved_review_asins,
)
from services.spapi import sp_api_configured, get_catalog_api
from services.storage_fees import extract_dimensions, calc_storage_fee

router = APIRouter(prefix="/analytics")


# --------------------------------------------------------------------------- #
# Raw-row file preview — NO header assumption.
# --------------------------------------------------------------------------- #

_PREVIEW_LIMIT = 25
MAX_PAGE_SIZE = 1000
MAX_TITLE_PAGES = 100   # explicit page count ceiling; 0 = unlimited (runner-bounded)


def _parse_raw_rows(filename: str, data: bytes, sheet_name: str = "") -> list[list[Any]]:
    return _parse_raw_rows_shared(filename, data, sheet_name=sheet_name)


@router.post("/preview")
async def analytics_preview(
    catalog_file: UploadFile = File(...),
    sheet_name: str = Form(""),
) -> dict[str, Any]:
    """Return the first N raw rows so the wizard can render a header-row picker.

    Also returns ``sheets`` — the list of sheet names for Excel workbooks so
    the wizard can offer a sheet-selector dropdown when the file has multiple tabs.
    Pass ``sheet_name`` to re-preview a specific sheet; omit to use the active sheet.
    """
    try:
        data = await read_upload_limited(catalog_file)
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"Could not read upload: {exc}")

    # Collect sheet names before parsing so we can surface them to the UI.
    filename = catalog_file.filename or ""
    sheets = _get_sheet_names(filename, data)

    # Resolve requested sheet — fall back to active if name not found.
    resolved_sheet = sheet_name if (sheet_name and sheet_name in sheets) else ""

    try:
        rows = _parse_raw_rows(filename, data, sheet_name=resolved_sheet)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"Could not parse file: {exc}")

    if not rows:
        raise HTTPException(status_code=400, detail="File appears to be empty")

    preview = rows[:_PREVIEW_LIMIT]
    normalised = []
    max_cols = 0
    for i, r in enumerate(preview):
        cells = [("" if c is None else str(c)) for c in r]
        max_cols = max(max_cols, len(cells))
        normalised.append({"row_number": i + 1, "cells": cells})
    for row in normalised:
        if len(row["cells"]) < max_cols:
            row["cells"] += [""] * (max_cols - len(row["cells"]))

    # The active sheet name: prefer the requested one, otherwise the first sheet.
    active_sheet = resolved_sheet or (sheets[0] if sheets else "")

    return {
        "filename": filename,
        "size": len(data),
        "total_rows": len(rows),
        "rows": normalised,
        "max_cols": max_cols,
        "sheets": sheets,
        "active_sheet": active_sheet,
    }


# --------------------------------------------------------------------------- #
# Status
# --------------------------------------------------------------------------- #


@router.get("/status")
def analytics_status() -> dict[str, Any]:
    return {"sp_api_configured": sp_api_configured()}


# --------------------------------------------------------------------------- #
# Runs list + detail
# --------------------------------------------------------------------------- #


def _decode_search_methods(value: Any) -> Any:
    if not value:
        return value
    try:
        return json.loads(value)
    except (ValueError, TypeError):
        return value


@router.get("/runs")
def list_analytics_runs() -> list[dict[str, Any]]:
    """Return every analytics run, newest first. Shape matches the UI's expectations."""
    with database._connect() as conn:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            """
            SELECT id, name, marketplace, search_methods, pages_per_title,
                   ai_clean_titles, total_catalog_items, total_candidates_found,
                   verified_count, review_count, not_approved_count,
                   status, progress_phase, progress_done, progress_total,
                   max_rank, min_rank, vetting_mode,
                   ai_check_status, ai_check_done, ai_check_total,
                   elig_check_status, elig_check_done, elig_check_total,
                   created_at, updated_at
            FROM analytics_runs
            ORDER BY created_at DESC
            """
        ).fetchall()

    out: list[dict[str, Any]] = []
    for r in rows:
        d = dict(r)
        d["search_methods"] = _decode_search_methods(d.get("search_methods"))
        out.append(d)
    return out


@router.get("/runs/{run_id}")
def get_run_detail(
    run_id: int,
    limit: int = 200,
    offset: int = 0,
    verdict: str = "",
    max_rank: int = 0,
    ai_verdict: str = "",
    search: str = "",
) -> dict[str, Any]:
    """
    Return a single run row + up to `limit` candidates.

    max_rank: when > 0, only return candidates whose sales_rank <= max_rank
              (or whose rank is unknown/null). Pass 0 to disable.
    ai_verdict: when set (approve|reject|uncertain), filter candidates by
                ai_verdict and return ai_filtered_count in the response.
    search: when set, filter candidates server-side by ASIN (exact, case-insensitive)
            or by partial match against the amazon title stored in data_json.
    """
    limit = clamp_int(limit, 1, MAX_PAGE_SIZE, 200)
    offset = max(0, int(offset or 0))
    search = (search or "").strip()

    with database._connect() as conn:
        conn.row_factory = sqlite3.Row
        run = conn.execute(
            "SELECT * FROM analytics_runs WHERE id=?", (run_id,),
        ).fetchone()
        if not run:
            raise HTTPException(status_code=404, detail=f"Run {run_id} not found")
        run_d = dict(run)
        run_d["search_methods"] = _decode_search_methods(run_d.get("search_methods"))

        # Catalog rows are only needed for the initial full-data load.
        # Verdict-filtered fetches skip them — the client already has them.
        if not verdict:
            rows = conn.execute(
                "SELECT row_idx, data_json FROM analytics_catalog_rows "
                "WHERE run_id=? ORDER BY row_idx",
                (run_id,),
            ).fetchall()
            catalog_rows = [json.loads(r["data_json"]) for r in rows]
        else:
            catalog_rows = []

        rank_clause   = "AND (sales_rank IS NULL OR sales_rank <= ?)" if max_rank > 0 else ""
        # Permanent pre-run rank exclusion: always hide items outside the range
        # set when the run was created, regardless of the toolbar filter.
        # Uses inline integers (not placeholders) since these come from our DB.
        # max_rank: null BSR passes (new/unsold listings may lack BSR).
        # min_rank: null BSR is also excluded — a floor means "only ranked items".
        _run_max_r = int(run_d.get("max_rank") or 0)
        _run_min_r = int(run_d.get("min_rank") or 0)
        _run_rank_clause = ""
        if _run_max_r > 0:
            _run_rank_clause += f" AND (sales_rank IS NULL OR sales_rank <= {_run_max_r})"
        if _run_min_r > 0:
            _run_rank_clause += f" AND sales_rank IS NOT NULL AND sales_rank >= {_run_min_r}"
        ai_clause     = "AND ai_verdict=?" if ai_verdict else ""
        # Server-side search: exact ASIN match first; otherwise title LIKE.
        # The amazon title lives inside data_json — use json_extract so we
        # don't need to load the full blob.  ASIN column is indexed so the
        # exact-match branch is fast even on 74k rows.
        search_clause = (
            "AND (UPPER(asin) = UPPER(?) OR "
            "json_extract(data_json,'$.amazon.title') LIKE ? ESCAPE '\\')"
        ) if search else ""
        select_cols = (
            "row_idx, asin, sources, confidence, verdict, amz_pack, "
            "sales_rank, review_status, ai_verdict, ai_reasoning, "
            "eligibility_status, storage_fee, storage_fee_peak, data_json"
        )

        def _base(extra: list) -> list:
            p: list = [run_id]
            if verdict:
                p.append(str(verdict))
            if max_rank > 0:
                p.append(int(max_rank))
            if ai_verdict:
                p.append(str(ai_verdict))
            if search:
                p.append(search)                        # exact ASIN match
                p.append(f"%{search.replace('%','\\%').replace('_','\\_')}%")  # LIKE title
            p.extend(extra)
            return p

        verdict_clause = "AND verdict=?" if verdict else ""
        where = f"WHERE run_id=? {verdict_clause} {_run_rank_clause} {rank_clause} {ai_clause} {search_clause}"

        # Total count for this filter combination (used for pagination when
        # ai_verdict is active so the client knows the real total).
        ai_filtered_count: int | None = None
        if ai_verdict:
            cnt_row = conn.execute(
                f"SELECT COUNT(*) AS cnt FROM analytics_candidates {where}",
                _base([]),
            ).fetchone()
            ai_filtered_count = cnt_row["cnt"] if cnt_row else 0

        # Per-verdict AI counts scoped to the current verdict tab (ignores the
        # ai_verdict filter so all three buckets are always returned).
        where_no_ai = f"WHERE run_id=? {verdict_clause} {_run_rank_clause} {rank_clause}"
        ai_counts_params: list = [run_id]
        if verdict:
            ai_counts_params.append(str(verdict))
        if max_rank > 0:
            ai_counts_params.append(int(max_rank))
        ai_count_rows = conn.execute(
            f"SELECT ai_verdict, COUNT(*) AS cnt FROM analytics_candidates "
            f"{where_no_ai} AND ai_verdict IS NOT NULL GROUP BY ai_verdict",
            ai_counts_params,
        ).fetchall()
        ai_verdict_counts = {r["ai_verdict"]: r["cnt"] for r in ai_count_rows}

        cand_rows = conn.execute(
            f"SELECT {select_cols} FROM analytics_candidates {where} "
            f"ORDER BY confidence DESC, row_idx, asin LIMIT ? OFFSET ?",
            _base([limit, offset]),
        ).fetchall()

    candidates: list[dict] = []
    for c in cand_rows:
        d = dict(c)
        try:
            d["sources"] = json.loads(d.get("sources") or "[]")
        except (ValueError, TypeError):
            d["sources"] = []
        try:
            d["data"] = json.loads(d.pop("data_json") or "{}")
        except (ValueError, TypeError):
            d["data"] = {}
        # Expose sales_rank at the top level for the UI; also pull from
        # data.amazon.sales_rank for rows that predate the column migration.
        if d.get("sales_rank") is None:
            amz = (d.get("data") or {}).get("amazon") or {}
            sr = amz.get("sales_rank")
            if sr is not None:
                d["sales_rank"] = sr
        candidates.append(d)

    # ASIN conflict map: query ALL candidates for this run (ignoring the current
    # verdict/rank filter) so we catch conflicts even when a filter is active.
    # An ASIN matched to more than one distinct catalog row means different
    # catalog entries resolved to the same Amazon product — almost always a
    # sign that the vendor catalog contains duplicate SKUs that slipped through
    # dedup (e.g. same item with different UPCs) or a genuine ambiguous match.
    with database._connect() as conn:
        conn.row_factory = sqlite3.Row
        all_asin_rows = conn.execute(
            f"SELECT DISTINCT asin, row_idx FROM analytics_candidates "
            f"WHERE run_id=? {_run_rank_clause}",
            (run_id,),
        ).fetchall()
    asin_to_rows: dict[str, list[int]] = {}
    for r in all_asin_rows:
        asin_to_rows.setdefault(r["asin"], []).append(r["row_idx"])
    asin_conflicts = {
        asin: sorted(set(row_idxs))
        for asin, row_idxs in asin_to_rows.items()
        if len(set(row_idxs)) > 1
    }

    # ASIN conflicts are surfaced to the user as an informational "⚠ Conflict"
    # badge (via `asin_conflicts` in the response) but no longer downgrade the
    # verdict.  Per the user's rule, a candidate scoring >= the auto-approve
    # threshold is a confirmed match and belongs in Approved even when the same
    # ASIN was matched to more than one catalog row — which in this catalog is
    # almost always a vendor-side duplicate (same product, two SKUs/UPCs).  The
    # badge still lets the user spot and manually discard a genuine ambiguous
    # match.  (Previously this loop capped conflicted verified items to Review
    # and persisted that downgrade, stranding 90%+ matches on the Review tab.)
    _to_conflict_cap: list[tuple[int, str]] = []  # retained for the refresh guard; cap disabled

    # Defensive promotion: a candidate stored as "not_approved" with confidence
    # >= REVIEW_FLOOR but no hard-reject flag and no valid BSR cap is in an
    # inconsistent state. Promote to "review" AND persist to DB so the verdict
    # tabs stay consistent (previously in-memory-only, which caused the item to
    # vanish from both "Not Approved" and "Review" tabs since the DB still said
    # not_approved but the response said review).
    _run_max_r = int(run_d.get("max_rank") or 0)
    _run_min_r = int(run_d.get("min_rank") or 0)
    _HARD_REJECT_FLAGS = frozenset([
        "size_mismatch", "gender_mismatch", "color_mismatch",
        "category_mismatch", "media_format_mismatch",
        "count_mismatch", "scent_mismatch", "shade_mismatch",
        "apparel_size_mismatch",
    ])
    _to_promote: list[tuple[int, str]] = []   # (row_idx, asin) pairs to fix in DB
    for c in candidates:
        if c.get("verdict") != "not_approved":
            continue
        if (c.get("review_status") or "").strip():
            continue  # respect manual decisions
        conf_val = float(c.get("confidence") or 0)
        if conf_val < REVIEW_FLOOR:
            continue
        sc = (c.get("data") or {}).get("scores") or {}
        if any(sc.get(f) for f in _HARD_REJECT_FLAGS):
            continue
        rank = c.get("sales_rank")
        bsr_capped = (
            (_run_max_r > 0 and rank is not None and int(rank) > _run_max_r) or
            (_run_min_r > 0 and rank is not None and int(rank) < _run_min_r)
        )
        if bsr_capped:
            continue
        c["verdict"] = "review"
        c["auto_promoted"] = True
        _to_promote.append((c["row_idx"], c["asin"]))

    # Persist promotions to DB in one batch so subsequent requests are consistent.
    if _to_promote:
        with database._LOCK, database._connect() as conn:
            conn.executemany(
                "UPDATE analytics_candidates SET verdict='review' "
                "WHERE run_id=? AND row_idx=? AND asin=? AND verdict='not_approved' "
                "AND (review_status IS NULL OR review_status = '')",
                [(run_id, row_idx, asin) for row_idx, asin in _to_promote],
            )
        # Re-sync run-level counts and patch run_d so the tab badges in THIS
        # response already reflect the promoted items (run_d was fetched before
        # the update, so review_count / not_approved_count are stale otherwise).
        with database._LOCK, database._connect() as conn:
            updated = _recompute_run_counts(conn, run_id)
        run_d.update(updated)

    # High-confidence promotion (run-wide): any candidate stored as "review"
    # with confidence >= AUTO_APPROVE and no hard-reject / pack mismatch / BSR
    # violation is a confirmed match.  The user's rule is explicit — confidence
    # >= 90 always belongs in Approved.  Such items are in Review only because an
    # earlier version conflict-capped them (now disabled) or they predate the
    # current scorer.  A single bulk UPDATE promotes them across the WHOLE run
    # (not just the current page) so the fix is complete in one load.  Manual /
    # AI decisions (non-empty review_status) are always respected.  Hard-reject
    # and pack flags live in data_json.scores, so they're checked via
    # json_extract; sales_rank is a real column for the BSR guard.
    _hard_reject_sql = " ".join(
        f"AND COALESCE(json_extract(data_json,'$.scores.{_f}'),0)=0"
        for _f in _HARD_REJECT_FLAGS
    )
    with database._LOCK, database._connect() as conn:
        _verify_cur = conn.execute(
            f"UPDATE analytics_candidates SET verdict='verified' "
            f"WHERE run_id=? AND verdict='review' "
            f"  AND (review_status IS NULL OR review_status='') "
            f"  AND confidence >= ? "
            f"  {_hard_reject_sql} "
            f"  AND COALESCE(json_extract(data_json,'$.scores.pack_mismatch'),0)=0 "
            f"  AND NOT (? > 0 AND sales_rank IS NOT NULL AND sales_rank > ?) "
            f"  AND NOT (? > 0 AND sales_rank IS NOT NULL AND sales_rank < ?)",
            (run_id, float(AUTO_APPROVE),
             _run_max_r, _run_max_r, _run_min_r, _run_min_r),
        )
        _to_verify = bool(_verify_cur.rowcount and _verify_cur.rowcount > 0)

    if _to_verify:
        # Reflect the promotion in the current page's in-memory candidates so
        # THIS response is consistent with the DB write above.
        for c in candidates:
            if c.get("verdict") != "review":
                continue
            if (c.get("review_status") or "").strip():
                continue
            if float(c.get("confidence") or 0) < AUTO_APPROVE:
                continue
            sc = (c.get("data") or {}).get("scores") or {}
            if any(sc.get(f) for f in _HARD_REJECT_FLAGS) or sc.get("pack_mismatch"):
                continue
            rank = c.get("sales_rank")
            if (_run_max_r > 0 and rank is not None and int(rank) > _run_max_r) \
               or (_run_min_r > 0 and rank is not None and int(rank) < _run_min_r):
                continue
            c["verdict"] = "verified"
            c["auto_verified"] = True
        with database._LOCK, database._connect() as conn:
            updated = _recompute_run_counts(conn, run_id)
        run_d.update(updated)

    # If conflict-cap / auto-promotion / high-conf promotion wrote new verdict
    # changes, the ai_verdict_counts computed inside the initial DB block is
    # stale (it was computed before those writes).  Refresh it so the scoped
    # counts the client stores in tabPageData match the actual DB state.
    if _to_conflict_cap or _to_promote or _to_verify:
        _ai_cnt_params: list = [run_id]
        _vc = f"AND verdict=?" if verdict else ""
        if verdict:
            _ai_cnt_params.append(str(verdict))
        if max_rank > 0:
            _ai_cnt_params.append(int(max_rank))
        _rc = f"AND (sales_rank IS NULL OR sales_rank <= ?)" if max_rank > 0 else ""
        with database._connect() as _rc_conn:
            _rc_conn.row_factory = sqlite3.Row
            _fresh_rows = _rc_conn.execute(
                f"SELECT ai_verdict, COUNT(*) AS cnt FROM analytics_candidates "
                f"WHERE run_id=? {_vc} {_rc} AND ai_verdict IS NOT NULL "
                f"GROUP BY ai_verdict",
                _ai_cnt_params,
            ).fetchall()
        ai_verdict_counts = {r["ai_verdict"]: r["cnt"] for r in _fresh_rows}

    resp: dict[str, Any] = {
        "run": run_d,
        "catalog_rows": catalog_rows,
        "candidates": candidates,
        "ai_verdict_counts": ai_verdict_counts,
        "asin_conflicts": asin_conflicts,
    }
    if ai_filtered_count is not None:
        resp["ai_filtered_count"] = ai_filtered_count
    return resp


# --------------------------------------------------------------------------- #
# POST /runs — the wizard endpoint
# --------------------------------------------------------------------------- #


def _parse_json_field(raw: str, default: Any) -> Any:
    if raw is None or raw == "":
        return default
    try:
        return json.loads(raw)
    except (ValueError, TypeError):
        return default


@router.post("/runs")
async def create_analytics_run(
    catalog_file: UploadFile = File(...),
    name: str = Form(...),
    marketplace: str = Form("US"),
    header_row: int = Form(...),
    mapping: str = Form(...),           # JSON-encoded {"upc": "3", ...}
    brand: str = Form(""),
    search_methods: str = Form(...),    # JSON-encoded list
    pages_per_title: int = Form(5),
    ai_clean_titles: bool = Form(False),
    max_rank: int = Form(0),
    min_rank: int = Form(0),
    vetting_mode: str = Form("cpg"),
    sheet_name: str = Form(""),
    passthrough_cols: str = Form(""),   # JSON array of vendor column header names to carry into export
) -> dict[str, Any]:
    """
    Create a new Analytics run.

    The wizard sends us a multipart POST with the raw catalog file plus
    every decision the user made in steps 1–4. We parse the file with
    their chosen header row + column mapping, persist the resulting
    source rows, and kick off a background thread that runs the
    3-tier SP-API search + vetting.

    ``sheet_name`` lets the caller specify which Excel sheet to read.
    If omitted or not found, the workbook's active sheet is used.

    Response shape:
        {"run_id": 17, "total_catalog_items": 142, "status": "Searching"}
    """
    try:
        data = await read_upload_limited(catalog_file)
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"Could not read upload: {exc}")

    mapping_dict = _parse_json_field(mapping, {})
    methods_list = _parse_json_field(search_methods, [])
    if not isinstance(mapping_dict, dict):
        raise HTTPException(status_code=400, detail="mapping must be a JSON object")
    if not isinstance(methods_list, list) or not methods_list:
        raise HTTPException(status_code=400, detail="search_methods must be a non-empty JSON array")

    # Parse the file into source rows using the wizard's picks.
    try:
        rows = parse_source_rows(
            filename=catalog_file.filename or "catalog.xlsx",
            data=data,
            header_row_idx=int(header_row),
            mapping=mapping_dict,
            brand=brand or "",
            sheet_name=sheet_name or "",
        )
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"Could not parse file: {exc}")

    if not rows:
        raise HTTPException(
            status_code=400,
            detail="No catalog rows found after applying your column mapping. "
                   "Check the header row and UPC / Item ID / Title columns.",
        )

    # ── Deduplicate catalog rows ─────────────────────────────────────────────
    # Vendor catalogs frequently repeat the same SKU on multiple lines.  A row is
    # treated as a duplicate only when it repeats a STRONG identity, keeping the
    # FIRST occurrence (its original row_idx is preserved so UI row numbers stay
    # meaningful).
    #
    # Key priority: UPC (the globally-unique per-product identifier) is preferred;
    # when an Item ID is also present, BOTH must match.  This is deliberate — some
    # vendor files leave the real "Item #" column blank and the wizard's Item ID
    # ends up mapped to a low-cardinality column (Manufacturer, Category).  Keying
    # on Item ID alone then collapses thousands of distinct products onto one row
    # per manufacturer.  Requiring the UPC (or UPC+ItemID together) prevents that,
    # while also never merging two genuinely different SKUs that share a case UPC.
    # Only when there is NO UPC do we fall back to Item ID — and even then we pair
    # it with the title so a low-cardinality Item ID can't collapse the file.
    def _dedup_key(r):
        norm_upc   = r.upc.strip()
        norm_id    = r.itemid.strip().upper().replace("-", "").replace(" ", "")
        norm_title = r.title.strip().lower()
        if norm_upc and norm_id:
            return ("upc+id", norm_upc, norm_id)
        if norm_upc:
            return ("upc", norm_upc)
        if norm_id and norm_title:
            return ("id+title", norm_id, norm_title)
        if norm_id:
            return ("id", norm_id)
        return ("title", norm_title)

    seen_keys: set = set()
    deduped = []
    for row in rows:
        k = _dedup_key(row)
        if k in seen_keys:
            continue
        seen_keys.add(k)
        deduped.append(row)

    duplicates_removed = len(rows) - len(deduped)
    rows = deduped

    # Determine brand column/mode for storage so the rescore modal can pre-populate.
    # When brand text is non-empty the wizard was in text mode; otherwise the brand
    # column name comes from mapping["brand"].
    saved_brand_col  = brand.strip() if brand.strip() else (mapping_dict.get("brand") or "")
    saved_brand_mode = "text" if brand.strip() else "col"

    # 0 = unlimited (walk every page Amazon returns); the runner bounds it with a
    # safety cap. Positive values are capped at MAX_TITLE_PAGES.
    pages_per_title = clamp_int(pages_per_title, 0, MAX_TITLE_PAGES, 5)

    # Validate passthrough_cols — must be a JSON array of strings if provided.
    pt_cols_raw = (passthrough_cols or "").strip()
    pt_cols_validated: list[str] = []
    if pt_cols_raw:
        try:
            parsed_pt = json.loads(pt_cols_raw)
            if isinstance(parsed_pt, list):
                pt_cols_validated = [str(c) for c in parsed_pt if c]
        except (ValueError, TypeError):
            pass  # ignore malformed input — passthrough is optional

    run_id = start_analytics_run(
        name=name.strip() or (catalog_file.filename or "Untitled run"),
        marketplace=marketplace or "US",
        search_methods=[str(m) for m in methods_list],
        pages_per_title=pages_per_title,
        ai_clean_titles=bool(ai_clean_titles),
        source_rows=rows,
        max_rank=int(max_rank),
        min_rank=int(min_rank),
        mode=vetting_mode if vetting_mode in ("cpg", "medical") else "cpg",
        brand_col=saved_brand_col,
        brand_mode=saved_brand_mode,
        passthrough_cols=json.dumps(pt_cols_validated) if pt_cols_validated else "",
    )

    # Persist the duplicate count so the run detail banner can show it.
    if duplicates_removed > 0:
        with database._LOCK, database._connect() as conn:
            conn.execute(
                "UPDATE analytics_runs SET duplicate_rows_removed=? WHERE id=?",
                (duplicates_removed, run_id),
            )

    return {
        "run_id": run_id,
        "total_catalog_items": len(rows),
        "duplicate_rows_removed": duplicates_removed,
        "status": "Searching" if sp_api_configured() else "Error",
        "sp_api_configured": sp_api_configured(),
    }


# --------------------------------------------------------------------------- #
# POST /runs/{run_id}/candidates/verdict — manual verdict override
# --------------------------------------------------------------------------- #

_VERDICT_VALUES = {"verified", "review", "not_approved"}


def _recompute_run_counts(conn: sqlite3.Connection, run_id: int) -> dict[str, int]:
    rows = conn.execute(
        "SELECT verdict, COUNT(*) AS n FROM analytics_candidates "
        "WHERE run_id=? GROUP BY verdict",
        (run_id,),
    ).fetchall()
    verified = review = not_approved = total = 0
    for r in rows:
        v = (r["verdict"] or "").lower()
        n = int(r["n"])
        total += n
        if v == "verified":
            verified += n
        elif v == "review":
            review += n
        elif v == "not_approved":
            not_approved += n
    conn.execute(
        "UPDATE analytics_runs SET total_candidates_found=?, "
        "  verified_count=?, review_count=?, not_approved_count=?, "
        "  updated_at=CURRENT_TIMESTAMP WHERE id=?",
        (total, verified, review, not_approved, run_id),
    )
    return {
        "total_candidates_found": total,
        "verified_count": verified,
        "review_count": review,
        "not_approved_count": not_approved,
    }


@router.post("/runs/{run_id}/candidates/verdict")
def update_candidate_verdict(
    run_id: int,
    body: dict = Body(...),
) -> dict[str, Any]:
    """
    Manually override the verdict on one candidate row. The wizard sets
    verdicts automatically from the confidence-score thresholds; this
    endpoint is the UI's "Promote" / "Reject" / "Approve" / "Discard"
    button — the same workflow the Vetting review modal uses.

    Body: {"row_idx": int, "asin": str, "verdict": "verified"|"review"|"not_approved",
           "review_status": str (optional)}
    """
    row_idx = body.get("row_idx")
    asin = body.get("asin")
    verdict = (body.get("verdict") or "").lower().strip()
    review_status = body.get("review_status") or ""

    if row_idx is None or asin is None:
        raise HTTPException(status_code=400, detail="row_idx and asin are required")
    if verdict not in _VERDICT_VALUES:
        raise HTTPException(
            status_code=400,
            detail=f"verdict must be one of {sorted(_VERDICT_VALUES)}",
        )

    with database._LOCK, database._connect() as conn:
        conn.row_factory = sqlite3.Row

        # Update the candidate row and also rewrite the embedded verdict
        # inside data_json so CSV exports / UI reloads stay consistent.
        existing = conn.execute(
            "SELECT data_json FROM analytics_candidates "
            "WHERE run_id=? AND row_idx=? AND asin=?",
            (int(run_id), int(row_idx), str(asin)),
        ).fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="Candidate not found")

        try:
            data = json.loads(existing["data_json"] or "{}")
        except (ValueError, TypeError):
            data = {}
        data["verdict"] = verdict
        if review_status:
            data["review_status"] = review_status

        conn.execute(
            "UPDATE analytics_candidates "
            "SET verdict=?, review_status=?, data_json=? "
            "WHERE run_id=? AND row_idx=? AND asin=?",
            (verdict, review_status, json.dumps(data),
             int(run_id), int(row_idx), str(asin)),
        )

        # Look up the UPC for this catalog row so we can persist the pair.
        cat_row = conn.execute(
            "SELECT data_json FROM analytics_catalog_rows "
            "WHERE run_id=? AND row_idx=?",
            (int(run_id), int(row_idx)),
        ).fetchone()
        upc = ""
        if cat_row:
            try:
                upc = json.loads(cat_row["data_json"] or "{}").get("upc") or ""
            except (ValueError, TypeError):
                pass

        counts = _recompute_run_counts(conn, int(run_id))

    # Persist to the global pair manager outside the write lock.
    if upc:
        if verdict == "not_approved":
            failed = list(data.get("scores", {}).keys()) if data.get("scores") else []
            conf_val = data.get("scores", {}).get("confidence_score")
            database.add_to_blacklist(upc, str(asin), conf_val, failed)
        elif verdict == "verified":
            database.save_verified(upc, str(asin), data, review_status)

    return {
        "ok": True,
        "run_id": int(run_id),
        "row_idx": int(row_idx),
        "asin": str(asin),
        "verdict": verdict,
        "review_status": review_status,
        "counts": counts,
    }


@router.post("/runs/{run_id}/candidates/bulk_verdict")
def bulk_update_candidate_verdict(
    run_id: int,
    body: dict = Body(...),
) -> dict[str, Any]:
    """
    Apply the same verdict to a list of (row_idx, asin) pairs. This powers
    the "Reject all" / "Approve all" / "Promote all" bulk actions in the
    Approved / Review / Not Approved tabs of the run detail view.

    Body: {"items": [{"row_idx": int, "asin": str}, ...],
           "verdict": "verified"|"review"|"not_approved",
           "review_status": str (optional)}
    """
    items = body.get("items") or []
    verdict = (body.get("verdict") or "").lower().strip()
    review_status = body.get("review_status") or ""

    if not isinstance(items, list) or not items:
        raise HTTPException(status_code=400, detail="items must be a non-empty list")
    if verdict not in _VERDICT_VALUES:
        raise HTTPException(
            status_code=400,
            detail=f"verdict must be one of {sorted(_VERDICT_VALUES)}",
        )

    updated = 0
    with database._LOCK, database._connect() as conn:
        conn.row_factory = sqlite3.Row
        for it in items:
            row_idx = it.get("row_idx")
            asin = it.get("asin")
            if row_idx is None or asin is None:
                continue
            existing = conn.execute(
                "SELECT data_json FROM analytics_candidates "
                "WHERE run_id=? AND row_idx=? AND asin=?",
                (int(run_id), int(row_idx), str(asin)),
            ).fetchone()
            if not existing:
                continue
            try:
                data = json.loads(existing["data_json"] or "{}")
            except (ValueError, TypeError):
                data = {}
            data["verdict"] = verdict
            if review_status:
                data["review_status"] = review_status
            conn.execute(
                "UPDATE analytics_candidates "
                "SET verdict=?, review_status=?, data_json=? "
                "WHERE run_id=? AND row_idx=? AND asin=?",
                (verdict, review_status, json.dumps(data),
                 int(run_id), int(row_idx), str(asin)),
            )
            updated += 1

        counts = _recompute_run_counts(conn, int(run_id))

    return {
        "ok": True,
        "run_id": int(run_id),
        "updated": updated,
        "verdict": verdict,
        "review_status": review_status,
        "counts": counts,
    }


@router.post("/runs/{run_id}/candidates/bulk_verdict_all")
def bulk_update_candidate_verdict_all(
    run_id: int,
    body: dict = Body(...),
) -> dict[str, Any]:
    """
    Apply a verdict to EVERY candidate currently in a given verdict bucket for the
    run — entirely server-side, so it never depends on which rows the client has
    loaded (the paginated detail view only holds one page per tab). Powers the
    per-tab "Approve/Reject/Promote All" buttons and the "Bulk approve" category
    modal.

    Body: {"from_verdict": "review"|"not_approved"|"verified",
           "to_verdict":   "verified"|"review"|"not_approved",
           "review_status": str (optional, e.g. "Manually Approved")}

    A non-empty review_status marks the change as a manual decision so the run
    detail's auto-promotion/conflict logic leaves it alone afterwards.
    """
    from_verdict = (body.get("from_verdict") or "").lower().strip()
    to_verdict = (body.get("to_verdict") or "").lower().strip()
    review_status = body.get("review_status") or ""
    if from_verdict not in _VERDICT_VALUES or to_verdict not in _VERDICT_VALUES:
        raise HTTPException(
            status_code=400,
            detail=f"from_verdict / to_verdict must be one of {sorted(_VERDICT_VALUES)}",
        )

    with database._LOCK, database._connect() as conn:
        conn.row_factory = sqlite3.Row
        # Single bulk UPDATE: flip the verdict column and mirror verdict +
        # review_status into data_json (via json_set) so the stored blob stays
        # consistent with the column the queries read.
        cur = conn.execute(
            "UPDATE analytics_candidates "
            "SET verdict=?, review_status=?, "
            "    data_json=json_set(data_json, '$.verdict', ?, '$.review_status', ?) "
            "WHERE run_id=? AND verdict=?",
            (to_verdict, review_status, to_verdict, review_status,
             int(run_id), from_verdict),
        )
        updated = cur.rowcount
        counts = _recompute_run_counts(conn, int(run_id))

    return {
        "ok": True,
        "run_id": int(run_id),
        "updated": updated,
        "from_verdict": from_verdict,
        "to_verdict": to_verdict,
        "review_status": review_status,
        "counts": counts,
    }


# --------------------------------------------------------------------------- #
# POST /runs/{run_id}/rescore — re-score candidates with a different title col
# --------------------------------------------------------------------------- #


@router.post("/runs/{run_id}/rescore")
def rescore_run(run_id: int, body: dict = Body(...)) -> dict[str, Any]:
    """
    Kick off a background re-score for an existing run using a different title
    column from the original catalog's raw data.

    Body: {"title_col": "Full Product Title"}
      - title_col: column header name from the raw dict stored in
        analytics_catalog_rows.data_json.raw.  Pass "" to use the
        originally-mapped title field.

    Returns immediately; progress is tracked via the run's progress_done/total
    fields (same mechanism as the search pipeline).
    """
    with database._connect() as conn:
        conn.row_factory = sqlite3.Row
        run = conn.execute(
            "SELECT id FROM analytics_runs WHERE id=?", (run_id,)
        ).fetchone()
        if not run:
            raise HTTPException(status_code=404, detail=f"Run {run_id} not found")

    title_col  = (body.get("title_col") or "").strip()
    brand_col  = (body.get("brand_col") or "").strip()
    brand_mode = (body.get("brand_mode") or "col").strip()
    max_rank   = int(body.get("max_rank") or 0)
    min_rank   = int(body.get("min_rank") or 0)
    start_rescore(int(run_id), title_col, brand_col, max_rank, min_rank, brand_mode)
    return {"ok": True, "started": True}


# --------------------------------------------------------------------------- #
# POST /runs/{run_id}/control — pause / stop / resume
# --------------------------------------------------------------------------- #

_ACTIVE_STATUSES   = {"Searching", "Pending", "Vetting"}
_PAUSABLE_STATUSES = {"Searching", "Vetting"}
_RESUMABLE_STATUSES = {"Paused", "Stopped"}
_STOPPABLE_STATUSES = _PAUSABLE_STATUSES | _RESUMABLE_STATUSES | {"Pending"}


@router.post("/runs/{run_id}/control")
def control_analytics_run(run_id: int, body: dict = Body(...)) -> dict[str, Any]:
    """
    Body: {"action": "pause" | "stop" | "resume"}

    pause  — signals the running pipeline to pause after its current step.
    stop   — signals the pipeline to stop and mark the run as Stopped.
    resume — restarts a Paused/Stopped run from where it left off.
    """
    action = (body.get("action") or "").lower().strip()
    if action not in ("pause", "stop", "resume"):
        raise HTTPException(status_code=400, detail="action must be pause, stop, or resume")

    with database._connect() as conn:
        conn.row_factory = sqlite3.Row
        run = conn.execute(
            "SELECT status FROM analytics_runs WHERE id=?", (run_id,)
        ).fetchone()
    if not run:
        raise HTTPException(status_code=404, detail=f"Run {run_id} not found")

    status = run["status"] or ""

    if action == "pause":
        if status not in _PAUSABLE_STATUSES:
            raise HTTPException(status_code=400, detail=f"Run is {status} — can only pause a running run")
        request_pause(run_id)
        # Write immediately so the UI reflects it even if the thread hasn't
        # reached its next control-check point yet.
        with database._LOCK, database._connect() as conn:
            conn.execute(
                "UPDATE analytics_runs SET status='Paused', "
                "progress_phase='Paused by user', updated_at=CURRENT_TIMESTAMP "
                "WHERE id=? AND status NOT IN ('Complete', 'Error', 'Stopped')",
                (run_id,),
            )
        return {"ok": True, "action": "pause", "run_id": run_id}

    if action == "stop":
        if status not in _STOPPABLE_STATUSES:
            raise HTTPException(status_code=400, detail=f"Run is {status} — nothing to stop")
        request_stop(run_id)
        # Write Stopped immediately so the UI reflects it even when the
        # background thread has died (e.g. after a server restart).
        with database._LOCK, database._connect() as conn:
            conn.execute(
                "UPDATE analytics_runs SET status='Stopped', "
                "progress_phase='Stopped by user', updated_at=CURRENT_TIMESTAMP "
                "WHERE id=? AND status NOT IN ('Complete', 'Error')",
                (run_id,),
            )
        return {"ok": True, "action": "stop", "run_id": run_id}

    if action == "resume":
        if status not in _RESUMABLE_STATUSES:
            raise HTTPException(status_code=400, detail=f"Run is {status} — can only resume a Paused or Stopped run")
        ok = resume_run(run_id)
        if not ok:
            raise HTTPException(status_code=400, detail="Could not resume run")
        return {"ok": True, "action": "resume", "run_id": run_id}


# --------------------------------------------------------------------------- #
# DELETE /runs/{run_id} — permanently remove a run and all its data
# --------------------------------------------------------------------------- #

@router.delete("/runs/{run_id}")
def delete_analytics_run(run_id: int) -> dict[str, Any]:
    request_stop(run_id)
    with database._LOCK, database._connect() as conn:
        run = conn.execute(
            "SELECT id FROM analytics_runs WHERE id=?", (run_id,)
        ).fetchone()
        if not run:
            raise HTTPException(status_code=404, detail=f"Run {run_id} not found")
        conn.execute("DELETE FROM analytics_runs WHERE id=?", (run_id,))
    return {"ok": True, "deleted": run_id}


# --------------------------------------------------------------------------- #
# AI check endpoints
# --------------------------------------------------------------------------- #


@router.get("/runs/{run_id}/ai_check/estimate")
def ai_check_estimate(run_id: int, verdict: str = "review") -> dict[str, Any]:
    """Return cost/time estimate for running AI check on candidates of this run.

    verdict may be a comma-separated list (e.g. "review,not_approved") or "all"/""
    to count every candidate.
    """
    with database._connect() as conn:
        conn.row_factory = sqlite3.Row
        verdicts = [v.strip() for v in verdict.split(",") if v.strip() and v.strip() != "all"]
        if verdicts:
            placeholders = ",".join("?" * len(verdicts))
            row = conn.execute(
                f"SELECT COUNT(*) AS cnt FROM analytics_candidates "
                f"WHERE run_id=? AND verdict IN ({placeholders})",
                [run_id] + verdicts,
            ).fetchone()
        else:
            row = conn.execute(
                "SELECT COUNT(*) AS cnt FROM analytics_candidates WHERE run_id=?",
                (run_id,),
            ).fetchone()
    count = row["cnt"] if row else 0
    est   = ai_check_cost(count)
    est["already_running"] = ai_check_running(run_id)
    return est


@router.post("/runs/{run_id}/ai_check")
def start_ai_check_endpoint(run_id: int, body: dict = Body(...)) -> dict[str, Any]:
    """Start AI check for candidates of this run in a background thread.

    body.verdict may be a comma-separated list (e.g. "review,not_approved"),
    "all", or "" to check every candidate.
    """
    with database._connect() as conn:
        conn.row_factory = sqlite3.Row
        run = conn.execute(
            "SELECT id FROM analytics_runs WHERE id=?", (run_id,)
        ).fetchone()
        if not run:
            raise HTTPException(status_code=404, detail=f"Run {run_id} not found")

    if ai_check_running(run_id):
        return {"ok": True, "already_running": True}

    # Reset the "applied" flag so the Apply button becomes available again after a re-check.
    with database._LOCK, database._connect() as conn:
        conn.execute(
            "UPDATE analytics_runs SET ai_decisions_applied=0, updated_at=CURRENT_TIMESTAMP WHERE id=?",
            (run_id,),
        )

    verdict_filter = (body.get("verdict") or "review").lower().strip()
    if verdict_filter == "all":
        verdict_filter = ""
    start_ai_check(run_id, verdict_filter)
    return {"ok": True, "started": True}


@router.post("/runs/{run_id}/ai_check/stop")
def stop_ai_check_endpoint(run_id: int) -> dict[str, Any]:
    """Signal the running AI check to stop immediately.

    If an in-memory thread is alive → set its stop Event (fast path).
    If the server was restarted and the thread is dead but the DB still says
    'Running' (orphan state) → update the DB directly so the UI clears.
    """
    running = ai_check_running(run_id)
    stop_ai_check(run_id)   # sets the Event; harmless if no thread is alive

    # Orphan guard: if no live thread owns this run, forcibly clear the DB status
    # so the UI doesn't show a permanent spinner.
    if not running:
        with database._connect() as conn:
            conn.row_factory = sqlite3.Row
            row = conn.execute(
                "SELECT ai_check_status FROM analytics_runs WHERE id=?", (run_id,)
            ).fetchone()
        if row and (row["ai_check_status"] or "").startswith("Running"):
            with database._LOCK, database._connect() as conn:
                conn.execute(
                    "UPDATE analytics_runs SET ai_check_status='Stopped' WHERE id=?",
                    (run_id,),
                )
            return {"ok": True, "was_running": False, "orphan_cleared": True}

    return {"ok": True, "was_running": running, "stopping": running}


# ── Eligibility (CAN_SELL/NEEDS_APPROVAL/RESTRICTED) + storage fee ─────────────
@router.get("/runs/{run_id}/eligibility/estimate")
def eligibility_estimate(run_id: int) -> dict[str, Any]:
    """How many DISTINCT Approved/Review ASINs would be checked (= live API calls)."""
    return {
        "unique_asins": approved_review_asins(run_id),
        "already_running": elig_check_running(run_id),
    }


@router.post("/runs/{run_id}/eligibility")
def start_eligibility_endpoint(run_id: int) -> dict[str, Any]:
    """Fetch eligibility + storage fee for this run's Approved/Review ASINs (background)."""
    with database._connect() as conn:
        conn.row_factory = sqlite3.Row
        run = conn.execute("SELECT id FROM analytics_runs WHERE id=?", (run_id,)).fetchone()
        if not run:
            raise HTTPException(status_code=404, detail=f"Run {run_id} not found")
    if elig_check_running(run_id):
        return {"ok": True, "already_running": True}
    start_eligibility_check(run_id)
    return {"ok": True, "started": True}


@router.post("/runs/{run_id}/eligibility/stop")
def stop_eligibility_endpoint(run_id: int) -> dict[str, Any]:
    """Signal the running eligibility check to stop; clear an orphaned status."""
    running = elig_check_running(run_id)
    stop_eligibility_check(run_id)
    if not running:
        with database._connect() as conn:
            conn.row_factory = sqlite3.Row
            row = conn.execute(
                "SELECT elig_check_status FROM analytics_runs WHERE id=?", (run_id,)
            ).fetchone()
        if row and (row["elig_check_status"] or "").startswith("Running"):
            with database._LOCK, database._connect() as conn:
                conn.execute(
                    "UPDATE analytics_runs SET elig_check_status='Stopped' WHERE id=?",
                    (run_id,),
                )
            return {"ok": True, "orphan_cleared": True}
    return {"ok": True, "was_running": running}


@router.post("/runs/{run_id}/ai_check/apply")
def apply_ai_decisions(run_id: int) -> dict[str, Any]:
    """Apply AI verdicts to candidate verdicts.

    ai_verdict='approve' → verdict='verified'
    ai_verdict='reject'  → verdict='not_approved'
    ai_verdict='uncertain' — left unchanged.
    Returns {approved: N, rejected: N, counts: {...}}.
    """
    with database._connect() as conn:
        conn.row_factory = sqlite3.Row
        run = conn.execute(
            "SELECT id FROM analytics_runs WHERE id=?", (run_id,)
        ).fetchone()
        if not run:
            raise HTTPException(status_code=404, detail=f"Run {run_id} not found")

    approved = rejected = 0
    with database._LOCK, database._connect() as conn:
        conn.row_factory = sqlite3.Row

        # Approve: ai_verdict='approve' → verdict='verified'
        # Set review_status='ai-accepted' so the ASIN-conflict cap and auto-promotion
        # guards (both check for non-empty review_status) never override this decision.
        rows_approve = conn.execute(
            "SELECT row_idx, asin, data_json FROM analytics_candidates "
            "WHERE run_id=? AND ai_verdict='approve'",
            (run_id,),
        ).fetchall()
        for r in rows_approve:
            try:
                data = json.loads(r["data_json"] or "{}")
            except (ValueError, TypeError):
                data = {}
            data["verdict"] = "verified"
            conn.execute(
                "UPDATE analytics_candidates "
                "SET verdict='verified', review_status='ai-accepted', data_json=? "
                "WHERE run_id=? AND row_idx=? AND asin=?",
                (json.dumps(data), run_id, r["row_idx"], r["asin"]),
            )
            approved += 1

        # Reject: ai_verdict='reject' → verdict='not_approved'
        # Set review_status='ai-rejected' so auto-promotion never reverses this decision.
        rows_reject = conn.execute(
            "SELECT row_idx, asin, data_json FROM analytics_candidates "
            "WHERE run_id=? AND ai_verdict='reject'",
            (run_id,),
        ).fetchall()
        for r in rows_reject:
            try:
                data = json.loads(r["data_json"] or "{}")
            except (ValueError, TypeError):
                data = {}
            data["verdict"] = "not_approved"
            conn.execute(
                "UPDATE analytics_candidates "
                "SET verdict='not_approved', review_status='ai-rejected', data_json=? "
                "WHERE run_id=? AND row_idx=? AND asin=?",
                (json.dumps(data), run_id, r["row_idx"], r["asin"]),
            )
            rejected += 1

        counts = _recompute_run_counts(conn, run_id)
        conn.execute(
            "UPDATE analytics_runs SET ai_decisions_applied=1, updated_at=CURRENT_TIMESTAMP WHERE id=?",
            (run_id,),
        )

    return {"ok": True, "approved": approved, "rejected": rejected, "counts": counts}


# --------------------------------------------------------------------------- #
# GET /runs/{run_id}/export — multi-sheet Excel download
# --------------------------------------------------------------------------- #

_HDR_FONT = Font(bold=True)


def _hdr_row(ws, titles: list[str]) -> list[WriteOnlyCell]:
    cells = []
    for t in titles:
        c = WriteOnlyCell(ws, value=t)
        c.font = _HDR_FONT
        cells.append(c)
    return cells


@router.get("/runs/{run_id}/export")
def export_analytics_run(
    run_id: int,
    min_rank: int = 0,
    max_rank: int = 0,
    skip_null_rank: bool = False,
):
    """
    Return a four-sheet Excel workbook filtered by the caller's rank window:
      Sheet 1 "Approved"     — verified candidates
      Sheet 2 "Review"       — review candidates
      Sheet 3 "Not Approved" — not_approved candidates
      Sheet 4 "Not Found"    — catalog rows with no Amazon match
    """
    with database._connect() as conn:
        conn.row_factory = sqlite3.Row
        run = conn.execute(
            "SELECT * FROM analytics_runs WHERE id=?", (run_id,)
        ).fetchone()
        if not run:
            raise HTTPException(status_code=404, detail=f"Run {run_id} not found")
        run_d = dict(run)

        # Use json_extract so we never load the full data_json blob into Python
        # for large runs (74k candidates × 3 KB = ~220 MB → OOM crash).
        # Only the two Amazon fields actually used in the export are extracted.
        cat_rows = conn.execute(
            """
            SELECT row_idx,
                   json_extract(data_json, '$.upc')    AS upc,
                   json_extract(data_json, '$.itemid') AS itemid,
                   json_extract(data_json, '$.title')  AS title,
                   json_extract(data_json, '$.brand')  AS brand,
                   json_extract(data_json, '$.raw')    AS raw_json
            FROM analytics_catalog_rows
            WHERE run_id=? ORDER BY row_idx
            """,
            (run_id,),
        ).fetchall()

        # Apply the run's permanent rank exclusion so BSR-filtered items are also
        # absent from the export (same logic as get_run_detail).
        _exp_run_max = int(run_d.get("max_rank") or 0)
        _exp_run_min = int(run_d.get("min_rank") or 0)
        _exp_rank_clause = ""
        if _exp_run_max > 0:
            _exp_rank_clause += f" AND (sales_rank IS NULL OR sales_rank <= {_exp_run_max})"
        if _exp_run_min > 0:
            _exp_rank_clause += f" AND sales_rank IS NOT NULL AND sales_rank >= {_exp_run_min}"

        cand_rows = conn.execute(
            f"""
            SELECT row_idx, asin, confidence, verdict, sales_rank, amz_pack,
                   eligibility_status, storage_fee, storage_fee_peak,
                   COALESCE(
                       json_extract(data_json, '$.amazon.title'), ''
                   ) AS amz_title,
                   COALESCE(
                       json_extract(data_json, '$.amazon.brand'),
                       json_extract(data_json, '$.amazon.manufacturer'),
                       ''
                   ) AS amz_brand
            FROM analytics_candidates
            WHERE run_id=? {_exp_rank_clause}
            ORDER BY confidence DESC
            """,
            (run_id,),
        ).fetchall()

    # Parse passthrough column list from the run metadata.
    pt_cols: list[str] = []
    try:
        pt_raw = run_d.get("passthrough_cols") or ""
        if pt_raw:
            parsed = json.loads(pt_raw)
            if isinstance(parsed, list):
                pt_cols = [str(c) for c in parsed if c]
    except (ValueError, TypeError):
        pass

    src_by_idx: dict[int, dict] = {}
    for r in cat_rows:
        d = dict(r)
        # Parse the raw dict once so passthrough lookups are O(1).
        try:
            d["_raw"] = json.loads(d.get("raw_json") or "{}") if pt_cols else {}
        except (ValueError, TypeError):
            d["_raw"] = {}
        src_by_idx[r["row_idx"]] = d

    found_idxs = {c["row_idx"] for c in cand_rows}

    def _rank_passes(sales_rank) -> bool:
        if sales_rank is None:
            return not skip_null_rank
        r = int(sales_rank)
        if min_rank > 0 and r < min_rank:
            return False
        if max_rank > 0 and r > max_rank:
            return False
        return True

    cand_hdr = [
        "Row", "Source UPC", "Source Item ID", "Source Title", "Source Brand",
        "ASIN", "Amazon Title", "Amazon Brand", "AMZ Pack", "BSR", "Confidence",
        "Approval Status", "Storage Fee/unit/mo", "Storage Fee/unit/mo (Q4 peak)",
    ] + pt_cols  # passthrough columns appended after fixed columns

    buckets: dict[str, list] = {"verified": [], "review": [], "not_approved": []}
    for c in cand_rows:
        if not _rank_passes(c["sales_rank"]):
            continue
        src = src_by_idx.get(c["row_idx"], {})
        raw = src.get("_raw") or {}
        pt_values = [raw.get(col, "") for col in pt_cols]
        row = safe_spreadsheet_row([
            (c["row_idx"] or 0) + 1,
            src.get("upc") or "",
            src.get("itemid") or "",
            src.get("title") or "",
            src.get("brand") or "",
            c["asin"] or "",
            c["amz_title"] or "",
            c["amz_brand"] or "",
            c["amz_pack"] if c["amz_pack"] else 1,   # no pack detected → 1 (single unit)
            c["sales_rank"] if c["sales_rank"] is not None else "",
            round(float(c["confidence"] or 0), 1),
            c["eligibility_status"] or "",           # blank until the eligibility check is run
            c["storage_fee"] if c["storage_fee"] is not None else "",
            c["storage_fee_peak"] if c["storage_fee_peak"] is not None else "",
        ] + pt_values)
        buckets.get((c["verdict"] or "").lower(), buckets["not_approved"]).append(row)

    wb = Workbook(write_only=True)

    for sheet_name, verdict_key in [
        ("Approved", "verified"),
        ("Review", "review"),
        ("Not Approved", "not_approved"),
    ]:
        ws = wb.create_sheet(sheet_name)
        ws.append(_hdr_row(ws, cand_hdr))
        for row in buckets[verdict_key]:
            ws.append(safe_spreadsheet_row(row))

    # Not Found sheet — unaffected by rank filter
    ws_nf = wb.create_sheet("Not Found")
    ws_nf.append(_hdr_row(ws_nf, [
        "Row", "Source UPC", "Source Item ID", "Source Title", "Source Brand",
    ] + pt_cols))
    for idx in sorted(src_by_idx.keys()):
        if idx in found_idxs:
            continue
        src = src_by_idx[idx]
        raw = src.get("_raw") or {}
        pt_values = [raw.get(col, "") for col in pt_cols]
        ws_nf.append(safe_spreadsheet_row([
            idx + 1,
            src.get("upc") or "",
            src.get("itemid") or "",
            src.get("title") or "",
            src.get("brand") or "",
        ] + pt_values))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # Strip non-ASCII characters so Content-Disposition filename is always valid.
    # Non-ASCII chars (e.g. Cyrillic) in unencoded filenames confuse browsers and
    # can cause the response to be saved as "export.txt" instead of the xlsx file.
    raw_name = str(run_d.get("name") or f"run-{run_id}").replace(" ", "_")
    ascii_name = "".join(
        c if c.isascii() and (c.isalnum() or c in "._-") else "_"
        for c in raw_name
    )
    safe_name = ascii_name.strip("_")[:120] or f"run-{run_id}"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{safe_name}_analytics.xlsx"'},
    )


# --------------------------------------------------------------------------- #
# Quick Search — transient single-item ASIN search, no DB writes
# --------------------------------------------------------------------------- #

@router.post("/quick-search")
def analytics_quick_search(body: dict = Body(...)) -> dict[str, Any]:
    """
    Synchronous single-item ASIN search with full confidence scoring.
    No DB writes — results are returned directly and not persisted.

    Body fields (all optional, at least one required):
      upc:          UPC / EAN barcode string
      itemid:       Item ID / MPN / SKU
      title:        Product title for keyword search
      brand:        Brand name applied to every candidate during scoring
      vetting_mode: "cpg" | "medical"  (default "cpg")
      max_rank:     int — BSR cap  (0 = disabled)
      min_rank:     int — BSR floor (0 = disabled)
    """
    if not sp_api_configured():
        raise HTTPException(status_code=503, detail="SP-API credentials not configured")

    upc      = str(body.get("upc")    or "").strip()
    itemid   = str(body.get("itemid") or "").strip()
    title    = str(body.get("title")  or "").strip()
    brand    = str(body.get("brand")  or "").strip()
    mode     = str(body.get("vetting_mode") or "cpg").lower()
    max_rank = int(body.get("max_rank") or 0)
    min_rank = int(body.get("min_rank") or 0)

    if not upc and not itemid and not title:
        raise HTTPException(
            status_code=400,
            detail="At least one of upc, itemid, or title is required",
        )

    # Zero-pad 11-digit UPCs → UPC-A
    if len(upc) == 11 and upc.isdigit():
        upc = "0" + upc

    row = _SourceRow(row_idx=0, upc=upc, itemid=itemid, title=title, brand=brand)
    source_dict = row.as_source_dict()
    api = get_catalog_api()

    # Collect (normalized_item, source_label) from all tiers.
    raw_candidates: list[tuple[dict, str]] = []

    if upc:
        # _tier1_upc returns (matched, kw_unverified) -- barcode-confirmed hits
        # vs. Pass-4 keyword-fallback hits whose barcode does NOT match the
        # searched UPC (often the same product under a different Amazon
        # UPC/ASIN). This caller was still unpacking it as a single dict (a
        # pre-tuple calling convention -- confirmed live 2026-09-22: crashed
        # with AttributeError on every UPC search), while the full Analytics
        # run's caller (services/analytics/runner.py ~line 1691) already
        # unpacks and tags both halves correctly. Matched here to that
        # standard: "UPC" gets full UPC credit in scoring below, "UPC-KW"
        # does not (still barcode/UPC-string-scoped search either way -- no
        # title/brand tier is involved, matching "only the UPC was searched").
        tier1_matched, tier1_kw = _tier1_upc(api, [row], run_id=0)
        for items in tier1_matched.values():
            for it in items:
                raw_candidates.append((it, "UPC"))
        for items in tier1_kw.values():
            for it in items:
                raw_candidates.append((it, "UPC-KW"))

    if itemid:
        tier2 = _tier2_itemid(api, [row], run_id=0)
        for items in tier2.values():
            for it in items:
                raw_candidates.append((it, "ItemID"))

    if title:
        tier3 = _tier3_title(api, [row], max_pages=3, run_id=0)
        for items in tier3.values():
            for it in items:
                raw_candidates.append((it, "Title"))

    # Deduplicate by ASIN, union sources.
    by_asin: dict[str, tuple[dict, list[str]]] = {}
    for item, src_label in raw_candidates:
        asin = (item.get("asin") or "").strip().upper()
        if not asin:
            continue
        if asin in by_asin:
            prev_item, prev_srcs = by_asin[asin]
            merged_srcs = list(dict.fromkeys(prev_srcs + [src_label]))
            by_asin[asin] = (prev_item, merged_srcs)
        else:
            by_asin[asin] = (item, [src_label])

    # Score each unique ASIN in-memory (mirrors _upsert_candidate logic, no DB).
    results: list[dict] = []
    for asin, (normalized, sources) in by_asin.items():
        scores = calculate_confidence(
            source_dict, normalized,
            upc_search_hit="UPC" in sources,
            mpn_search_hit="ItemID" in sources,
            mode=mode,
        )
        conf = scores["confidence_score"]

        # Category check (medical mode only).
        amz_bsr_cat = categorize("", normalized.get("sales_rank_category") or "")
        category_mismatch = False
        if mode == "medical" and amz_bsr_cat != UNKNOWN:
            dist = category_distance(MEDICAL, amz_bsr_cat)
            category_mismatch = dist >= 5.0
        if category_mismatch:
            scores["category_mismatch"] = True

        # Media-format hard-reject (both modes).
        media_mismatch = _is_media_format(normalized.get("title") or "")
        if media_mismatch:
            scores["media_format_mismatch"] = True

        hard_reject = (
            scores.get("size_mismatch") or scores.get("gender_mismatch")
            or scores.get("color_mismatch") or category_mismatch or media_mismatch
        )
        if hard_reject:
            verdict = "not_approved"
        elif conf >= AUTO_APPROVE:
            verdict = "verified"
        elif conf >= REVIEW_FLOOR or scores.get("pack_mismatch"):
            verdict = "review"
        elif conf >= MIN_CONFIDENCE:
            verdict = "review"
        else:
            verdict = "not_approved"

        sales_rank = normalized.get("sales_rank")
        if max_rank > 0 and sales_rank is not None and int(sales_rank) > max_rank:
            verdict = "not_approved"
        # min_rank excludes unranked items (null BSR) — same logic as the run filter.
        if min_rank > 0 and (sales_rank is None or int(sales_rank) < min_rank):
            verdict = "not_approved"

        # Dimensions, storage fee, and list price all come from the SAME raw
        # SP-API item this search already fetched (normalize_amazon_item keeps
        # the full response as `_raw`) -- no extra API call, same approach
        # services/analytics/eligibility_check.py already uses for a full
        # Analytics run's storage-fee enrichment (its own docstring: "computed
        # from the Amazon dimensions ALREADY stored ... no extra API call").
        # Quick Search fetches the identical data but was discarding it.
        raw_attrs = ((normalized.get("_raw") or {}).get("attributes") or {})
        dims = extract_dimensions(raw_attrs)
        storage_offpeak = storage_peak = None
        if all(dims.get(k) is not None for k in ("length_cm", "width_cm", "height_cm", "weight_g")):
            storage_offpeak, storage_peak = calc_storage_fee(
                dims["length_cm"], dims["width_cm"], dims["height_cm"], dims["weight_g"],
            )
        list_price = None
        lp_list = raw_attrs.get("list_price") or []
        if isinstance(lp_list, list) and lp_list and isinstance(lp_list[0], dict):
            list_price = lp_list[0].get("value")

        results.append({
            "asin":              asin,
            "title":             normalized.get("title") or "",
            "brand":             normalized.get("brand") or "",
            "manufacturer":      normalized.get("manufacturer") or "",
            "upc":               normalized.get("upc") or "",
            "ean":               normalized.get("ean") or "",
            "gtin":              normalized.get("gtin") or "",
            "mpn":               normalized.get("mpn") or "",
            "sales_rank":        sales_rank,
            "sales_rank_category": normalized.get("sales_rank_category") or "",
            "confidence":        round(conf, 1),
            "verdict":           verdict,
            "sources":           sources,
            "scores":            scores,
            "list_price":        list_price,
            "length_in":         dims["length_in"],
            "width_in":          dims["width_in"],
            "height_in":         dims["height_in"],
            "weight_lb":         dims["weight_lb"],
            "storage_fee_offpeak": storage_offpeak,
            "storage_fee_peak":    storage_peak,
        })

    # Context filter: when the user provides a brand or title, keep only
    # candidates where ALL significant words (>2 chars) from that context
    # appear in the Amazon brand + title.  This prevents the search from
    # returning unrelated products that happened to carry the same MPN/UPC.
    # Example: title="smith & nephew" → only keep listings that contain
    # both "smith" and "nephew" somewhere in the Amazon brand or title.
    context_text = (brand or title or "").strip().lower()
    if context_text:
        import re as _re
        context_words = [w for w in _re.findall(r'\b[a-z]{3,}\b', context_text)]
        if context_words:
            def _context_matches(r: dict) -> bool:
                haystack = (
                    (r.get("brand") or "") + " " +
                    (r.get("manufacturer") or "") + " " +
                    (r.get("title") or "")
                ).lower()
                return all(w in haystack for w in context_words)
            results = [r for r in results if _context_matches(r)]

    results.sort(key=lambda x: x["confidence"], reverse=True)
    return {"candidates": results, "total": len(results)}
